import csv
import io
import json
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.utils import timezone
from django.urls import reverse
from django.core.exceptions import ImproperlyConfigured
from openpyxl import Workbook, load_workbook
from PIL import Image

from gather_rsvp.database import database_config_from_url

from .models import (
    CrewMember,
    Event,
    EventMedia,
    Expense,
    Guest,
    GuestField,
    WhatsAppCampaign,
    WhatsAppCampaignRecipient,
)
from .views import guest_whatsapp_number


@override_settings(LOCAL_DEMO_MODE=False)
class OrganizerAccessTests(TestCase):
    def setUp(self):
        self.event = Event.objects.create(name="Test Event")
        self.guest = Guest.objects.create(
            event=self.event,
            name="Riya Test",
            phone="+911234567890",
        )
        self.user = get_user_model().objects.create_user(
            username="organizer",
            password="a-strong-test-password",
            is_staff=True,
        )

    def test_database_configuration_supports_secure_postgres_url_and_local_sqlite(self):
        local = database_config_from_url("", Path("local.sqlite3"))
        self.assertEqual(local["ENGINE"], "django.db.backends.sqlite3")
        self.assertEqual(local["NAME"], Path("local.sqlite3"))

        hosted = database_config_from_url(
            "postgresql://event_user:p%40ss@db.example.test:5432/nimoli"
            "?sslmode=verify-full",
            Path("unused.sqlite3"),
        )
        self.assertEqual(hosted["ENGINE"], "django.db.backends.postgresql")
        self.assertEqual(hosted["NAME"], "nimoli")
        self.assertEqual(hosted["USER"], "event_user")
        self.assertEqual(hosted["PASSWORD"], "p@ss")
        self.assertEqual(hosted["HOST"], "db.example.test")
        self.assertEqual(hosted["PORT"], "5432")
        self.assertEqual(hosted["OPTIONS"], {"sslmode": "verify-full"})

        with self.assertRaises(ImproperlyConfigured):
            database_config_from_url("mysql://db.example.test/nimoli", Path("local.sqlite3"))

    def test_guest_data_pages_require_authentication(self):
        for url in (
            reverse("dashboard"),
            reverse("export_guests"),
            reverse("event_media_file", args=[1]),
        ):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertRedirects(
                    response,
                    f"{reverse('login')}?next={url}",
                    fetch_redirect_response=False,
                )

    def test_guest_mutations_require_authentication(self):
        add_response = self.client.post(
            reverse("add_guest"),
            {
                "name": "New guest",
                "phone": "9999999999",
                "country_code": "+91",
                "guest_type": Guest.GuestType.LOCAL_STAY,
                "side_tag": Guest.SideTag.BRIDE,
                "message_consent": "yes",
            },
        )
        update_response = self.client.post(
            reverse("update_guest", args=[self.guest.pk]),
            {
                "rsvp_status": Guest.RSVPStatus.CONFIRMED,
                "hotel_status": Guest.HotelStatus.ONE_NIGHT,
                "whatsapp_opt_in": "yes",
            },
        )
        upload_response = self.client.post(
            reverse("upload_event_media"),
            {},
        )
        self.assertEqual(add_response.status_code, 302)
        self.assertEqual(update_response.status_code, 302)
        self.assertEqual(upload_response.status_code, 302)
        self.assertEqual(Guest.objects.count(), 1)
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.rsvp_status, Guest.RSVPStatus.PENDING)

    def test_authenticated_non_organizer_cannot_access_guest_data(self):
        non_staff = get_user_model().objects.create_user(
            username="attendee",
            password="a-strong-test-password",
        )
        self.client.force_login(non_staff)
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 403)
        self.assertEqual(self.client.get(reverse("export_guests")).status_code, 403)

    def test_authenticated_organizer_can_add_guest_and_record_opt_in(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("add_guest"),
            {
                "name": "New guest",
                "phone": "9999999999",
                "country_code": "+91",
                "guest_type": Guest.GuestType.LOCAL_STAY,
                "side_tag": Guest.SideTag.BRIDE,
                "message_consent": "yes",
            },
        )
        guest = Guest.objects.get(phone="9999999999")
        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        self.assertTrue(guest.whatsapp_opt_in)
        self.assertIsNotNone(guest.whatsapp_opt_in_at)

    def test_opt_out_and_invalid_status_updates(self):
        self.client.force_login(self.user)
        self.guest.whatsapp_opt_in = True
        self.guest.save(update_fields=["whatsapp_opt_in"])
        update_url = reverse("update_guest", args=[self.guest.pk])
        response = self.client.post(
            update_url,
            {
                "rsvp_status": Guest.RSVPStatus.CONFIRMED,
                "hotel_status": Guest.HotelStatus.TWO_NIGHTS,
                "whatsapp_opt_in": "no",
            },
        )
        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        self.guest.refresh_from_db()
        self.assertFalse(self.guest.whatsapp_opt_in)
        self.assertIsNotNone(self.guest.whatsapp_opted_out_at)
        self.assertEqual(self.guest.rsvp_status, Guest.RSVPStatus.CONFIRMED)

        self.client.post(
            update_url,
            {
                "rsvp_status": "unexpected",
                "hotel_status": Guest.HotelStatus.ONE_NIGHT,
                "whatsapp_opt_in": "yes",
            },
        )
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.rsvp_status, Guest.RSVPStatus.CONFIRMED)
        self.assertFalse(self.guest.whatsapp_opt_in)

    @override_settings(PUBLIC_BASE_URL="https://events.example.test")
    def test_csv_export_requires_login_and_escapes_formula_cells(self):
        self.client.force_login(self.user)
        self.guest.name = "=HYPERLINK(\"https://example.invalid\")"
        self.guest.save(update_fields=["name"])
        response = self.client.get(reverse("export_guests"))
        self.assertEqual(response.status_code, 200)
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(
            rows[0],
            [
                "Group ID",
                "First Name",
                "Last Name",
                "Country Code",
                "Phone Number",
                "Invited Dates",
                "Category",
                "Pax",
                "Guest Type",
                "Tag",
                "Sub Category",
                "Invite Name",
                "Salutation",
            ],
        )
        self.assertEqual(
            rows[1][rows[0].index("First Name")],
            "'=HYPERLINK(\"https://example.invalid\")",
        )

    def test_csv_export_matches_the_guest_template_columns(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("export_guests"))
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(len(rows[0]), 13)
        self.assertEqual(len(rows[1]), 13)

    def test_csv_import_uses_group_serial_and_allows_blank_family_contact(self):
        self.client.force_login(self.user)
        upload = SimpleUploadedFile(
            "guests.csv",
            b"Serial Number,First Name,Country Code,Contact Number,Guest Type,Tag\r\n"
            b"2,New Guest,+91,9999999999,Outstation,Bride's side\r\n"
            b"2,Family Member,+91,,Outstation,Bride's side\r\n",
            content_type="text/csv",
        )
        response = self.client.post(
            reverse("import_guests"),
            {"guest_file": upload},
            follow=True,
        )
        self.assertEqual(
            Guest.objects.count(),
            3,
            [str(message) for message in response.context["messages"]],
        )
        new_guest = Guest.objects.get(name="New Guest")
        family_member = Guest.objects.get(name="Family Member")
        self.assertFalse(new_guest.whatsapp_opt_in)
        self.assertEqual(new_guest.group_serial, family_member.group_serial)
        self.assertEqual(new_guest.country_code, "+91")
        self.assertEqual(new_guest.phone, "9999999999")
        self.assertEqual(family_member.phone, "")

    def test_csv_import_rejects_invalid_batch_without_partial_changes(self):
        self.client.force_login(self.user)
        upload = SimpleUploadedFile(
            "guests.csv",
            (
                "First Name,Last Name,Tag\r\n"
                "Valid,Guest,Bride Side\r\n"
                f"{'X' * 151},Too Long,Groom Side\r\n"
            ).encode(),
            content_type="text/csv",
        )
        response = self.client.post(
            reverse("import_guests"), {"guest_file": upload}, follow=True
        )
        self.assertEqual(Guest.objects.count(), 1)
        self.assertTrue(
            any("exceeds 150 characters" in str(message)
                for message in response.context["messages"])
        )

    def test_csv_import_accepts_partial_rows_and_assigns_group_ids(self):
        self.client.force_login(self.user)
        upload = SimpleUploadedFile(
            "partial-guests.csv",
            b"Last Name,Category\r\nSharma,Cousin\r\n,Friend\r\n",
            content_type="text/csv",
        )

        response = self.client.post(
            reverse("import_guests"), {"guest_file": upload}, follow=True
        )

        self.assertEqual(response.status_code, 200)
        guests = list(
            Guest.objects.filter(event=self.event, group_serial__in=[1, 2]).order_by(
                "group_serial"
            )
        )
        self.assertEqual(len(guests), 2)
        self.assertEqual([guest.group_serial for guest in guests], [1, 2])
        self.assertEqual(guests[0].name, "")
        self.assertEqual(guests[0].last_name, "Sharma")
        self.assertEqual(guests[0].category, "Cousin")
        self.assertIsNone(guests[0].pax)
        self.assertEqual(guests[0].side_tag, "")
        self.assertEqual(guests[0].guest_type, "")
        self.assertEqual(guests[0].country_code, "")
        self.assertEqual(guests[0].phone, "")

    def test_csv_import_preserves_unrecognized_partial_values(self):
        self.client.force_login(self.user)
        upload = SimpleUploadedFile(
            "partial-values.csv",
            b"Group ID,Tag,Pax,Invited Dates,Sub Category\r\n"
            b"not-set,Other,Several,Next Friday,Family\r\n",
            content_type="text/csv",
        )

        response = self.client.post(
            reverse("import_guests"), {"guest_file": upload}, follow=True
        )

        self.assertEqual(response.status_code, 200)
        guest = Guest.objects.get(event=self.event, group_serial=1)
        saved_values = {
            field.label: guest.custom_data.get(str(field.pk))
            for field in GuestField.objects.all()
        }
        self.assertEqual(guest.side_tag, "")
        self.assertIsNone(guest.pax)
        self.assertEqual(saved_values["Provided Group ID"], "not-set")
        self.assertEqual(saved_values["Provided Tag"], "Other")
        self.assertEqual(saved_values["Provided Pax"], "Several")
        self.assertEqual(saved_values["Provided Invited Dates"], "Next Friday")
        self.assertEqual(saved_values["Provided Sub Category"], "Family")

    def test_excel_import_saves_guest_invitation_fields(self):
        workbook = Workbook()
        bride_sheet = workbook.active
        bride_sheet.title = "Bride Side"
        groom_sheet = workbook.create_sheet("Groom Side")
        headers = [
            "Serial Number",
            "First Name",
            "Country Code",
            "Contact Number",
            "Guest Type",
            "Invited Date",
            "Category / Relation",
            "Pax",
            "Sub Category",
            "Invite Name",
            "Salutation",
        ]
        bride_sheet.append(headers)
        groom_sheet.append(headers)
        bride_sheet.append(
            [
                22,
                "New Guest",
                "+91",
                "9876543210",
                "International",
                date(2026, 12, 1),
                "Cousin",
                2,
                "Couple",
                "Aarav and Mira",
                "Dr.",
            ]
        )
        groom_sheet.append(
            [22, "Groom Side Guest", "+1", "2025550123", "International"]
        )
        content = io.BytesIO()
        workbook.save(content)
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("import_guests"),
            {
                "guest_file": SimpleUploadedFile(
                    "guests.xlsx",
                    content.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            Guest.objects.filter(name="New Guest").count(),
            1,
            [str(message) for message in response.context["messages"]],
        )
        guest = Guest.objects.get(name="New Guest")
        self.assertEqual(guest.group_serial, 22)
        self.assertEqual(guest.phone, "9876543210")
        self.assertEqual(guest.country_code, "+91")
        self.assertEqual(guest.invited_date, date(2026, 12, 1))
        self.assertEqual(guest.category, "Cousin")
        self.assertEqual(guest.pax, 2)
        self.assertEqual(guest.guest_type, Guest.GuestType.INTERNATIONAL)
        self.assertEqual(guest.side_tag, Guest.SideTag.BRIDE)
        self.assertEqual(guest.sub_category, Guest.SubCategory.COUPLE)
        self.assertEqual(guest.invite_name, "Aarav and Mira")
        self.assertEqual(guest.salutation, "Dr.")
        self.assertFalse(guest.whatsapp_opt_in)
        groom_guest = Guest.objects.get(name="Groom Side Guest")
        self.assertEqual(groom_guest.group_serial, guest.group_serial)
        self.assertEqual(groom_guest.side_tag, Guest.SideTag.GROOM)

    def test_organizer_can_record_guest_consent_for_campaigns(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("update_guest", args=[self.guest.pk]),
            {
                "rsvp_status": Guest.RSVPStatus.PENDING,
                "hotel_status": Guest.HotelStatus.NOT_SELECTED,
                "whatsapp_opt_in": "yes",
                "country_code": "+91",
                "phone": "9876543210",
                "group_serial": "1",
            },
        )

        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        self.guest.refresh_from_db()
        self.assertTrue(self.guest.whatsapp_opt_in)
        self.assertIsNotNone(self.guest.whatsapp_opt_in_at)
        self.assertIsNone(self.guest.whatsapp_opted_out_at)
        self.assertContains(self.client.get(reverse("dashboard")), "Opted in")

    def test_excel_export_contains_guest_data_and_rsvp_submission(self):
        self.client.force_login(self.user)
        self.guest.group_serial = 5
        self.guest.category = "Friend"
        self.guest.side_tag = Guest.SideTag.GROOM
        self.guest.last_name = "Guest"
        self.guest.country_code = "+91"
        self.guest.phone = "9876543210"
        self.guest.invited_date = date(2026, 10, 1)
        self.guest.pax = 3
        self.guest.save()
        Guest.objects.create(
            event=self.event,
            name="A Guest",
            group_serial=5,
            side_tag=Guest.SideTag.GROOM,
        )
        Guest.objects.create(
            event=self.event,
            name="Earlier Group",
            group_serial=2,
            side_tag=Guest.SideTag.GROOM,
        )

        response = self.client.get(reverse("export_excel"))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        sheet = workbook["Groom Side"]
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, f"A1:{sheet.cell(1, sheet.max_column).column_letter}4")
        workbook.close()
        header_indexes = {header: index for index, header in enumerate(rows[0])}
        self.assertEqual(rows[1][header_indexes["First Name"]], "Earlier Group")
        self.assertEqual(rows[2][header_indexes["First Name"]], "A Guest")
        guest_row = rows[3]
        self.assertEqual(guest_row[header_indexes["Group ID"]], 5)
        self.assertEqual(guest_row[header_indexes["Last Name"]], "Guest")
        self.assertEqual(guest_row[header_indexes["Country Code"]], "+91")
        self.assertEqual(guest_row[header_indexes["Phone Number"]], "9876543210")
        self.assertEqual(guest_row[header_indexes["Tag"]], "Groom Side")
        self.assertEqual(guest_row[header_indexes["Category"]], "Friend")
        self.assertEqual(guest_row[header_indexes["Pax"]], 3)
        self.assertEqual(guest_row[header_indexes["Invited Dates"]], "2026-10-01")
        self.assertEqual(
            list(rows[0]),
            [
                "Group ID",
                "First Name",
                "Last Name",
                "Country Code",
                "Phone Number",
                "Invited Dates",
                "Category",
                "Pax",
                "Guest Type",
                "Tag",
                "Sub Category",
                "Invite Name",
                "Salutation",
            ],
        )

    def test_guest_excel_template_uses_eventiq_columns_and_dropdowns(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("guest_template"))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), read_only=False)
        self.assertEqual(
            workbook.sheetnames,
            ["Instructions", "Guest List", "Lists"],
        )
        instructions = "\n".join(
            str(cell.value)
            for row in workbook["Instructions"].iter_rows()
            for cell in row
            if cell.value
        )
        self.assertIn("do not need to create separate Bride Side and Groom Side sheets", instructions)
        self.assertIn("sort or filter", instructions)
        self.assertIn("Invited Dates", instructions)
        sheet = workbook["Guest List"]
        self.assertEqual(
            [sheet.cell(row=1, column=index).value for index in range(1, 14)],
            [
                "Group ID",
                "First Name",
                "Last Name",
                "Country Code",
                "Phone Number",
                "Invited Dates",
                "Category",
                "Pax",
                "Guest Type",
                "Tag",
                "Sub Category",
                "Invite Name",
                "Salutation",
            ],
        )
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:M2001")
        validations = {
            validation.sqref.__str__(): validation.formula1
            for validation in sheet.data_validations.dataValidation
        }
        self.assertEqual(
            validations["I2:I2001"],
            '"Outstation,Local Stay,Local Non-Stay,International"',
        )
        self.assertTrue(
            next(
                validation
                for validation in sheet.data_validations.dataValidation
                if str(validation.sqref) == "I2:I2001"
            ).allow_blank
        )
        self.assertEqual(validations["J2:J2001"], '"Bride Side,Groom Side"')
        self.assertEqual(validations["D2:D2001"], "=CountryCodes")
        self.assertEqual(workbook["Lists"].sheet_state, "hidden")
        self.assertIn("+91", [row[0] for row in workbook["Lists"].iter_rows(min_row=2, values_only=True)])
        workbook.close()

    def test_excel_import_accepts_reordered_columns_and_common_headers(self):
        workbook = Workbook()
        bride_sheet = workbook.active
        bride_sheet.title = "Bride Side"
        groom_sheet = workbook.create_sheet("Groom Side")
        headers = [
            "Guest Type",
            "Phone Number",
            "First Name",
            "Last Name",
            "Group ID",
            "Country Dial Code",
            "Dietary Preference",
        ]
        bride_sheet.append(headers)
        groom_sheet.append(headers)
        bride_sheet.append(
            ["Local Stay", "9876543210", "Reordered", "Guest", 12, "+91", "Vegetarian"]
        )
        content = io.BytesIO()
        workbook.save(content)
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("import_guests"),
            {
                "guest_file": SimpleUploadedFile(
                    "guests.xlsx",
                    content.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            follow=True,
        )

        guest = Guest.objects.get(name="Reordered")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(guest.group_serial, 12)
        self.assertEqual(guest.country_code, "+91")
        self.assertEqual(guest.phone, "9876543210")
        self.assertEqual(guest.last_name, "Guest")
        custom_field = GuestField.objects.get(label="Dietary Preference")
        self.assertEqual(guest.custom_data[str(custom_field.pk)], "Vegetarian")

    def test_eventiq_workbook_import_uses_exact_headers_and_ignores_read_me(self):
        workbook = Workbook()
        guests_sheet = workbook.active
        guests_sheet.title = "Guests"
        guests_sheet.append(
            [
                "Group ID",
                "First Name",
                "Last Name",
                "Country Code",
                "Phone Number",
                "Invited Dates",
                "Category",
                "Pax",
                "Guest Type",
                "Tag",
                "Sub Category",
                "Invite Name",
                "Salutation",
            ]
        )
        guests_sheet.append(
            [
                31,
                "Riya",
                "Sharma",
                "+91",
                "9876543210",
                date(2026, 12, 1),
                "Cousin",
                2,
                "Local Stay",
                "Bride Side",
                "Couple",
                "Sharma Family",
                "Dr.",
            ]
        )
        read_me = workbook.create_sheet("Read Me")
        read_me.append(["This sheet is informational"])
        read_me.append(["It must not be interpreted as guest data"])
        content = io.BytesIO()
        workbook.save(content)
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("import_guests"),
            {
                "guest_file": SimpleUploadedFile(
                    "EventIQ_Guest_Import_Template.xlsx",
                    content.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        guest = Guest.objects.get(name="Riya", last_name="Sharma")
        self.assertEqual(guest.group_serial, 31)
        self.assertEqual(guest.phone, "9876543210")
        self.assertEqual(guest.invited_date, date(2026, 12, 1))
        self.assertEqual(guest.side_tag, Guest.SideTag.BRIDE)
        self.assertEqual(guest.invite_name, "Sharma Family")

    def test_excel_import_routes_a_single_sheet_by_guest_side_column(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Guest List"
        sheet.append(
            [
                "Guest Type",
                "Guest Side",
                "Contact Number",
                "First Name",
                "Serial Number",
                "Country Code",
                "Invited Dates",
                "Category",
                "Pax",
                "Sub Category",
                "Invite Name",
                "Salutation",
                "Dietary Preference",
            ]
        )
        sheet.append(
            [
                "Local Stay", "Groom Side", "9876543210", "Groom Guest", 7,
                "+91", date(2026, 12, 1), "Cousin", 2, "Couple",
                "Groom Family", "Dr.", "Vegetarian",
            ]
        )
        sheet.append(
            [
                "International", "Bride", "2025550123", "Bride Guest", 7,
                "+1", date(2026, 12, 2), "Friend", 1, "Single",
                "Bride Guest", "Ms.", "Vegan",
            ]
        )
        content = io.BytesIO()
        workbook.save(content)
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("import_guests"),
            {
                "guest_file": SimpleUploadedFile(
                    "guest-list.xlsx",
                    content.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        bride_guest = Guest.objects.get(name="Bride Guest")
        groom_guest = Guest.objects.get(name="Groom Guest")
        self.assertEqual(bride_guest.side_tag, Guest.SideTag.BRIDE)
        self.assertEqual(groom_guest.side_tag, Guest.SideTag.GROOM)
        self.assertEqual(bride_guest.group_serial, 7)
        self.assertEqual(groom_guest.group_serial, 7)
        self.assertEqual(bride_guest.invited_date, date(2026, 12, 2))
        self.assertEqual(groom_guest.category, "Cousin")
        self.assertEqual(groom_guest.pax, 2)
        self.assertEqual(groom_guest.sub_category, Guest.SubCategory.COUPLE)
        self.assertEqual(groom_guest.invite_name, "Groom Family")
        self.assertEqual(groom_guest.salutation, "Dr.")
        dietary_field = GuestField.objects.get(label="Dietary Preference")
        self.assertEqual(
            bride_guest.custom_data[str(dietary_field.pk)],
            "Vegan",
        )
        self.assertEqual(
            groom_guest.custom_data[str(dietary_field.pk)],
            "Vegetarian",
        )

    def test_excel_import_accepts_a_single_untagged_sheet_with_selected_side(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Guests"
        sheet.append(
            ["Serial Number", "First Name", "Country Code", "Contact Number", "Guest Type"]
        )
        sheet.append([21, "One Sheet Guest", "+91", "9876543210", "Local Stay"])
        content = io.BytesIO()
        workbook.save(content)
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("import_guests"),
            {
                "guest_file": SimpleUploadedFile(
                    "one-sheet.xlsx",
                    content.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
                "default_side": Guest.SideTag.BRIDE,
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        guest = Guest.objects.get(name="One Sheet Guest")
        self.assertEqual(guest.side_tag, Guest.SideTag.BRIDE)

    def test_excel_import_accepts_guest_without_whatsapp_contact(self):
        workbook = Workbook()
        bride_sheet = workbook.active
        bride_sheet.title = "Bride Side"
        groom_sheet = workbook.create_sheet("Groom Side")
        headers = [
            "Serial Number",
            "First Name",
            "Country Code",
            "Contact Number",
            "Guest Type",
        ]
        bride_sheet.append(headers)
        groom_sheet.append(headers)
        bride_sheet.append([1, "No Contact", "+91", "", "Local Stay"])
        content = io.BytesIO()
        workbook.save(content)
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("import_guests"),
            {
                "guest_file": SimpleUploadedFile(
                    "guests.xlsx",
                    content.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            follow=True,
        )

        self.assertEqual(Guest.objects.count(), 2)
        guest = Guest.objects.get(name="No Contact")
        self.assertEqual(guest.group_serial, 1)
        self.assertEqual(guest.phone, "")
        self.assertFalse(guest.whatsapp_opt_in)

    def test_manual_guest_can_be_created_with_partial_details(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("add_guest"),
            {
                "last_name": "Sharma",
                "country_code": "",
                "phone": "",
                "guest_type": "",
                "side_tag": "",
            },
        )

        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        guest = Guest.objects.get(event=self.event, last_name="Sharma")
        self.assertEqual(guest.name, "")
        self.assertEqual(guest.group_serial, 1)
        self.assertEqual(guest.side_tag, "")
        self.assertEqual(guest.guest_type, "")

    def test_manual_guest_contact_is_stored_without_messaging_consent(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("add_guest"),
            {
                "name": "No Consent",
                "country_code": "+91",
                "phone": "9876543210",
                "guest_type": "",
                "side_tag": "",
            },
        )

        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        guest = Guest.objects.get(event=self.event, name="No Consent")
        self.assertEqual(guest.phone, "9876543210")
        self.assertFalse(guest.whatsapp_opt_in)
        self.assertIsNone(guest.whatsapp_opt_in_at)

    def test_guest_rsvp_link_records_attendance_and_hotel_preferences(self):
        dietary_field = GuestField.objects.create(label="Dietary preference")
        url = reverse("guest_rsvp", args=[self.guest.rsvp_token])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Cache-Control"], "no-store")
        self.assertEqual(response["Referrer-Policy"], "no-referrer")
        self.assertEqual(response["X-Robots-Tag"], "noindex, nofollow, noarchive")
        self.assertContains(response, self.guest.name)
        self.assertEqual(response.status_code, 200)
        response = self.client.post(
            url,
            {
                "attendance": "yes",
                "room_preference": Guest.RoomPreference.DOUBLE,
                "hotel_check_in": "2026-11-12",
                "hotel_check_out": "2026-11-14",
                "pax": "2",
                f"custom_{dietary_field.pk}": "Vegetarian",
            },
        )
        self.assertRedirects(response, url, fetch_redirect_response=False)
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.rsvp_status, Guest.RSVPStatus.CONFIRMED)
        self.assertEqual(self.guest.room_preference, Guest.RoomPreference.DOUBLE)
        self.assertEqual(self.guest.hotel_check_in, date(2026, 11, 12))
        self.assertEqual(self.guest.hotel_check_out, date(2026, 11, 14))
        self.assertEqual(self.guest.pax, 2)
        self.assertEqual(
            self.guest.custom_data[str(dietary_field.pk)],
            "Vegetarian",
        )

    @override_settings(
        PUBLIC_BASE_URL="https://events.example.test",
        WHATSAPP_PHONE_NUMBER_ID="1234567890",
        WHATSAPP_ACCESS_TOKEN="test-token",
        WHATSAPP_API_VERSION="v22.0",
    )
    def test_whatsapp_campaign_queue_only_includes_opted_in_valid_guests(self):
        self.client.force_login(self.user)
        self.guest.whatsapp_opt_in = True
        self.guest.whatsapp_opt_in_at = timezone.now()
        self.guest.save(update_fields=["whatsapp_opt_in", "whatsapp_opt_in_at"])
        other_guest = Guest.objects.create(
            event=self.event,
            name="Not opted in",
            phone="+919999999999",
        )

        response = self.client.post(
            reverse("queue_whatsapp_campaign"),
            {"template_name": "event_invitation", "language_code": "en"},
        )

        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        campaign = WhatsAppCampaign.objects.get()
        self.assertEqual(campaign.status, WhatsAppCampaign.Status.QUEUED)
        self.assertEqual(campaign.recipient_count, 1)
        self.assertEqual(
            list(campaign.recipients.values_list("guest_id", flat=True)),
            [self.guest.pk],
        )
        self.assertFalse(campaign.recipients.filter(guest=other_guest).exists())

    @override_settings(
        PUBLIC_BASE_URL="https://events.example.test",
        WHATSAPP_PHONE_NUMBER_ID="1234567890",
        WHATSAPP_ACCESS_TOKEN="test-token",
        WHATSAPP_API_VERSION="v22.0",
        WHATSAPP_SEND_INTERVAL=0,
    )
    @patch(
        "guests.management.commands.send_whatsapp_campaigns.urlopen",
        return_value=io.BytesIO(b'{"messages":[{"id":"wamid.test-message"}]}'),
    )
    def test_sender_command_sends_individual_template_and_records_provider_id(self, urlopen):
        self.guest.whatsapp_opt_in = True
        self.guest.country_code = "+91"
        self.guest.save(update_fields=["whatsapp_opt_in", "country_code"])
        campaign = WhatsAppCampaign.objects.create(
            event=self.event,
            template_name="event_invitation",
            language_code="en",
            recipient_count=1,
        )
        recipient = WhatsAppCampaignRecipient.objects.create(
            campaign=campaign,
            guest=self.guest,
        )

        call_command("send_whatsapp_campaigns", confirm=True, stdout=io.StringIO())

        recipient.refresh_from_db()
        campaign.refresh_from_db()
        self.assertEqual(recipient.status, WhatsAppCampaignRecipient.Status.SENT)
        self.assertEqual(recipient.provider_message_id, "wamid.test-message")
        self.assertEqual(campaign.status, WhatsAppCampaign.Status.COMPLETED)
        request = urlopen.call_args.args[0]
        payload = json.loads(request.data.decode("utf-8"))
        self.assertEqual(payload["recipient_type"], "individual")
        self.assertEqual(payload["to"], "911234567890")
        self.assertEqual(
            payload["template"]["components"][0]["parameters"][0]["text"],
            "Riya Test",
        )

    @override_settings(
        PUBLIC_BASE_URL="https://events.example.test",
        WHATSAPP_PHONE_NUMBER_ID="1234567890",
        WHATSAPP_ACCESS_TOKEN="test-token",
        WHATSAPP_API_VERSION="v22.0",
        WHATSAPP_SEND_INTERVAL=0,
    )
    @patch("guests.management.commands.send_whatsapp_campaigns.urlopen")
    def test_sender_dry_run_does_not_send_and_send_requires_confirmation(self, urlopen):
        campaign = WhatsAppCampaign.objects.create(
            event=self.event,
            template_name="event_invitation",
            language_code="en",
            recipient_count=1,
        )
        WhatsAppCampaignRecipient.objects.create(
            campaign=campaign,
            guest=self.guest,
        )

        output = io.StringIO()
        call_command("send_whatsapp_campaigns", dry_run=True, stdout=output)
        self.assertIn("No messages sent", output.getvalue())
        urlopen.assert_not_called()

        with self.assertRaises(CommandError):
            call_command("send_whatsapp_campaigns", stdout=io.StringIO())
        urlopen.assert_not_called()

    def test_guest_rsvp_rejects_invalid_date_order_and_allows_decline(self):
        url = reverse("guest_rsvp", args=[self.guest.rsvp_token])
        self.client.post(
            url,
            {
                "attendance": "yes",
                "room_preference": Guest.RoomPreference.SINGLE,
                "hotel_check_in": "2026-11-14",
                "hotel_check_out": "2026-11-12",
            },
        )
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.rsvp_status, Guest.RSVPStatus.PENDING)

        self.client.post(url, {"attendance": "no"})
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.rsvp_status, Guest.RSVPStatus.DECLINED)
        self.assertIsNone(self.guest.hotel_check_in)

    def test_login_page_renders(self):
        response = self.client.get(reverse("login"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Welcome back")

    def test_pwa_manifest_worker_and_offline_page(self):
        manifest = self.client.get(reverse("web_manifest"))
        self.assertEqual(manifest.status_code, 200)
        self.assertEqual(manifest["Content-Type"], "application/manifest+json")
        self.assertEqual(manifest.json()["display"], "standalone")
        self.assertEqual(manifest.json()["name"], "Nimoli Bhandari Events")
        self.assertEqual(manifest.json()["short_name"], "Nimoli Bhandari")
        self.assertEqual(
            {icon["sizes"] for icon in manifest.json()["icons"]},
            {"192x192", "512x512"},
        )

        worker = self.client.get(reverse("service_worker"))
        self.assertEqual(worker.status_code, 200)
        self.assertEqual(worker["Service-Worker-Allowed"], "/")
        self.assertContains(worker, "request.mode === \"navigate\"")

        offline = self.client.get(reverse("offline"))
        self.assertEqual(offline.status_code, 200)
        self.assertContains(offline, "Guest information is not available offline")

    @override_settings(LOCAL_DEMO_MODE=True)
    def test_local_demo_opens_dashboard_and_export_without_sign_in(self):
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Local demo · sign-in bypassed")
        self.assertEqual(self.client.get(reverse("export_guests")).status_code, 200)

    @override_settings(LOCAL_DEMO_MODE=True)
    def test_local_demo_does_not_bypass_for_non_loopback_requests(self):
        response = self.client.get(
            reverse("dashboard"),
            REMOTE_ADDR="192.168.1.20",
        )
        self.assertRedirects(
            response,
            f"{reverse('login')}?next={reverse('dashboard')}",
            fetch_redirect_response=False,
        )

    def test_staff_dashboard_renders_guest_controls(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.guest.name)
        self.assertContains(response, "Log out")
        self.assertContains(response, "Bride’s side")
        self.assertContains(response, "Groom’s side")
        self.assertContains(response, "Extra details")
        self.assertContains(response, "Import Excel/CSV")
        self.assertContains(response, "Export Excel")
        self.assertContains(response, "<th>Group ID</th><th>First Name</th>")
        self.assertContains(response, "<th>Last Name</th><th>Country Code</th>")
        self.assertContains(response, "<th>Phone Number</th><th>Invited Dates</th>")
        self.assertContains(response, "<th>Guest Type</th><th>Tag</th>")
        self.assertContains(response, "<th>Invite Name</th><th>Salutation</th>")
        self.assertContains(response, "WhatsApp consent")
        self.assertContains(response, "Only select opted in after the guest has agreed")
        self.assertContains(response, 'aria-label="Guest list sheets"')
        self.assertContains(response, 'src="/static/guests/app-icon-192.png"')
        self.assertContains(response, 'aria-labelledby="event-media-heading"')
        self.assertContains(response, "Private to organizers")

    def test_event_media_upload_gallery_and_private_file_access(self):
        self.client.force_login(self.user)
        image_data = io.BytesIO()
        Image.new("RGB", (2, 2), color="green").save(image_data, format="PNG")
        upload = SimpleUploadedFile(
            "welcome.png",
            image_data.getvalue(),
            content_type="image/png",
        )

        with tempfile.TemporaryDirectory() as media_root, override_settings(
            MEDIA_ROOT=media_root
        ):
            response = self.client.post(
                reverse("upload_event_media"),
                {"media_file": upload, "title": "Welcome night"},
                follow=True,
            )

            self.assertEqual(response.status_code, 200)
            media = EventMedia.objects.get(event=self.event)
            self.assertEqual(media.title, "Welcome night")
            self.assertEqual(media.media_type, EventMedia.MediaType.IMAGE)
            self.assertEqual(media.content_type, "image/png")
            media_path = Path(media.file.path)
            self.assertTrue(media_path.is_file())
            self.assertContains(
                response,
                reverse("event_media_file", args=[media.pk]),
            )

            file_response = self.client.get(
                reverse("event_media_file", args=[media.pk])
            )
            self.assertEqual(file_response.status_code, 200)
            self.assertEqual(file_response["Content-Type"], "image/png")
            self.assertEqual(file_response["Cache-Control"], "private, no-store")
            self.assertEqual(file_response["X-Content-Type-Options"], "nosniff")
            self.assertEqual(
                b"".join(file_response.streaming_content),
                image_data.getvalue(),
            )
            file_response.close()

            delete_response = self.client.post(
                reverse("delete_event_media", args=[media.pk])
            )
            self.assertEqual(delete_response.status_code, 302)
            self.assertFalse(EventMedia.objects.filter(pk=media.pk).exists())
            self.assertFalse(media_path.exists())

    def test_event_media_rejects_fake_image_upload(self):
        self.client.force_login(self.user)
        upload = SimpleUploadedFile(
            "fake.png",
            b"not an image",
            content_type="image/png",
        )

        with tempfile.TemporaryDirectory() as media_root, override_settings(
            MEDIA_ROOT=media_root
        ):
            response = self.client.post(
                reverse("upload_event_media"),
                {"media_file": upload},
                follow=True,
            )

            self.assertEqual(response.status_code, 200)
            self.assertFalse(EventMedia.objects.exists())
            self.assertContains(response, "photo file is damaged")

    def test_event_media_upload_accepts_webm_signature(self):
        self.client.force_login(self.user)
        upload = SimpleUploadedFile(
            "short.webm",
            b"\x1a\x45\xdf\xa3" + b"\x00" * 20,
            content_type="video/webm",
        )

        with tempfile.TemporaryDirectory() as media_root, override_settings(
            MEDIA_ROOT=media_root
        ):
            response = self.client.post(
                reverse("upload_event_media"),
                {"media_file": upload, "title": "Ceremony"},
            )

            self.assertEqual(response.status_code, 302)
            media = EventMedia.objects.get(event=self.event)
            self.assertEqual(media.media_type, EventMedia.MediaType.VIDEO)
            self.assertEqual(media.content_type, "video/webm")

    def test_event_media_files_are_scoped_to_active_event(self):
        self.client.force_login(self.user)
        image_data = io.BytesIO()
        Image.new("RGB", (2, 2), color="blue").save(image_data, format="PNG")

        with tempfile.TemporaryDirectory() as media_root, override_settings(
            MEDIA_ROOT=media_root
        ):
            media = EventMedia.objects.create(
                event=self.event,
                file=SimpleUploadedFile("private.png", image_data.getvalue()),
                media_type=EventMedia.MediaType.IMAGE,
                content_type="image/png",
            )
            other_event = Event.objects.create(name="Different Event")
            self.client.post(reverse("select_event", args=[other_event.pk]))

            response = self.client.get(
                reverse("event_media_file", args=[media.pk])
            )

            self.assertEqual(response.status_code, 404)

    def test_guest_metadata_can_be_edited_and_side_filter_separates_lists(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("update_guest", args=[self.guest.pk]),
            {
                "rsvp_status": Guest.RSVPStatus.PENDING,
                "hotel_status": Guest.HotelStatus.NOT_SELECTED,
                "whatsapp_opt_in": "no",
                "country_code": "+91",
                "invited_date": "2026-12-01",
                "category": "Cousin",
                "last_name": "Updated",
                "pax": "2",
                "guest_type": Guest.GuestType.OUTSTATION,
                "side_tag": Guest.SideTag.BRIDE,
                "sub_category": Guest.SubCategory.COUPLE,
                "invite_name": "Riya and family",
                "salutation": "Ms.",
            },
        )
        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.country_code, "+91")
        self.assertEqual(self.guest.invited_date, date(2026, 12, 1))
        self.assertEqual(self.guest.category, "Cousin")
        self.assertEqual(self.guest.last_name, "Updated")
        self.assertEqual(self.guest.pax, 2)
        self.assertEqual(self.guest.guest_type, Guest.GuestType.OUTSTATION)
        self.assertEqual(self.guest.side_tag, Guest.SideTag.BRIDE)
        self.assertEqual(self.guest.sub_category, Guest.SubCategory.COUPLE)
        self.assertEqual(self.guest.invite_name, "Riya and family")
        self.assertEqual(self.guest.salutation, "Ms.")

        groom_guest = Guest.objects.create(
            event=self.event,
            name="Groom Guest",
            phone="+919999999999",
            side_tag=Guest.SideTag.GROOM,
        )
        response = self.client.get(reverse("dashboard"), {"side": Guest.SideTag.BRIDE})
        self.assertContains(response, "Bride Side")
        self.assertContains(response, self.guest.name)
        self.assertNotContains(response, groom_guest.name)

    def test_event_creation_and_switching_scopes_guest_list(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("event_list"),
            {"name": "Second Event", "venue": "Mumbai"},
        )
        second_event = Event.objects.get(name="Second Event")
        self.assertRedirects(response, reverse("event_list"), fetch_redirect_response=False)
        event_list_response = self.client.get(reverse("event_list"))
        self.assertContains(event_list_response, self.event.name)
        self.assertContains(event_list_response, second_event.name)
        self.assertEqual(Event.objects.count(), 2)
        self.assertEqual(self.client.session["active_event_id"], second_event.pk)
        self.assertNotContains(self.client.get(reverse("dashboard")), self.guest.name)

        self.client.post(reverse("select_event", args=[self.event.pk]))
        self.assertContains(self.client.get(reverse("dashboard")), self.guest.name)

    @override_settings(PUBLIC_BASE_URL="https://events.example.test")
    def test_opted_in_guest_gets_private_whatsapp_rsvp_link(self):
        self.guest.whatsapp_opt_in = True
        self.guest.save(update_fields=["whatsapp_opt_in"])
        self.client.force_login(self.user)

        response = self.client.get(reverse("dashboard"))

        guest = response.context["guests"][0]
        link = guest.whatsapp_invite_link
        parsed_link = urlparse(link)
        self.assertEqual(parsed_link.scheme, "https")
        self.assertEqual(parsed_link.netloc, "wa.me")
        self.assertEqual(parsed_link.path, "/911234567890")
        message = parse_qs(parsed_link.query)["text"][0]
        self.assertIn(
            f"https://events.example.test{reverse('guest_rsvp', args=[self.guest.rsvp_token])}",
            message,
        )
        self.assertIn("Reply STOP", message)

    @override_settings(PUBLIC_BASE_URL="http://events.example.test")
    def test_whatsapp_invite_link_requires_https(self):
        self.guest.whatsapp_opt_in = True
        self.guest.save(update_fields=["whatsapp_opt_in"])
        self.client.force_login(self.user)

        response = self.client.get(reverse("dashboard"))
        self.assertIsNone(response.context["guests"][0].whatsapp_invite_link)

    @override_settings(PUBLIC_BASE_URL="https://events.example.test")
    def test_whatsapp_invite_link_requires_opt_in_and_international_number(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"))
        self.assertIsNone(response.context["guests"][0].whatsapp_invite_link)

        self.guest.whatsapp_opt_in = True
        self.guest.phone = "9876543210"
        self.guest.save(update_fields=["whatsapp_opt_in", "phone"])
        response = self.client.get(reverse("dashboard"))
        self.assertIsNone(response.context["guests"][0].whatsapp_invite_link)

    def test_same_phone_can_be_added_to_a_different_event(self):
        self.client.force_login(self.user)
        other_event = Event.objects.create(name="Another Event")
        self.client.post(reverse("select_event", args=[other_event.pk]))
        response = self.client.post(
            reverse("add_guest"),
            {
                "name": "Riya at second event",
                "phone": "1234567890",
                "country_code": "+91",
                "guest_type": Guest.GuestType.OUTSTATION,
                "side_tag": Guest.SideTag.BRIDE,
                "message_consent": "yes",
            },
        )
        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        self.assertEqual(
            sum(
                guest_whatsapp_number(guest) == "911234567890"
                for guest in Guest.objects.all()
            ),
            2,
        )

    def test_manual_guest_without_contact_can_join_contacted_family_group(self):
        self.client.force_login(self.user)
        contact_guest = Guest.objects.create(
            event=self.event,
            name="Family contact",
            phone="9876543210",
            country_code="+91",
            side_tag=Guest.SideTag.BRIDE,
            group_serial=8,
        )

        response = self.client.post(
            reverse("add_guest"),
            {
                "name": "Family member",
                "phone": "",
                "country_code": "",
                "group_serial": "8",
                "guest_type": Guest.GuestType.LOCAL_STAY,
                "side_tag": Guest.SideTag.BRIDE,
            },
        )

        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        family_member = Guest.objects.get(name="Family member")
        self.assertEqual(family_member.group_serial, contact_guest.group_serial)
        self.assertFalse(family_member.whatsapp_opt_in)

    def test_expense_add_total_validation_and_event_scoped_delete(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("expense_list"),
            {"title": "Venue deposit", "category": "Venue", "amount": "1200.50"},
        )
        self.assertRedirects(response, reverse("expense_list"), fetch_redirect_response=False)
        expense = Expense.objects.get(event=self.event)
        self.assertEqual(expense.amount, Decimal("1200.50"))
        self.assertContains(self.client.get(reverse("expense_list")), "1200.50")

        self.client.post(
            reverse("expense_list"),
            {"title": "Invalid precision", "amount": "1.001"},
        )
        self.assertEqual(Expense.objects.count(), 1)

        other_event = Event.objects.create(name="Other Event")
        other_expense = Expense.objects.create(
            event=other_event,
            title="Other venue",
            amount=Decimal("1.00"),
        )
        response = self.client.post(reverse("delete_expense", args=[other_expense.pk]))
        self.assertEqual(response.status_code, 404)
        self.client.post(reverse("delete_expense", args=[expense.pk]))
        self.assertFalse(Expense.objects.filter(pk=expense.pk).exists())

    def test_crew_member_create_validation_and_removal(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("crew_list"),
            {
                "name": "Asha Crew",
                "role": "Coordinator",
                "phone": "+919876543210",
                "email": "invalid",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(CrewMember.objects.count(), 0)

        self.client.post(
            reverse("crew_list"),
            {
                "name": "Asha Crew",
                "role": "Coordinator",
                "phone": "+919876543210",
                "email": "asha@example.test",
            },
        )
        member = CrewMember.objects.get(event=self.event)
        self.assertContains(self.client.get(reverse("crew_list")), "Asha Crew")
        self.client.post(reverse("delete_crew_member", args=[member.pk]))
        self.assertFalse(CrewMember.objects.filter(pk=member.pk).exists())

    def test_organizer_can_add_custom_column_update_guest_and_export_value(self):
        self.client.force_login(self.user)
        self.client.post(reverse("add_guest_field"), {"label": "Dietary preference"})
        field = GuestField.objects.get(label="Dietary preference")
        update_url = reverse("update_guest", args=[self.guest.pk])
        self.client.post(
            update_url,
            {
                "rsvp_status": Guest.RSVPStatus.PENDING,
                "hotel_status": Guest.HotelStatus.NOT_SELECTED,
                "whatsapp_opt_in": "no",
                f"custom_{field.pk}": "Vegetarian",
            },
        )
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.custom_data[str(field.pk)], "Vegetarian")
        response = self.client.get(reverse("dashboard"))
        self.assertContains(response, "Dietary preference")
        self.assertContains(response, "Vegetarian")

        response = self.client.get(reverse("export_guests"))
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        custom_index = rows[0].index("Dietary preference")
        self.assertEqual(rows[1][custom_index], "Vegetarian")

    def test_removing_custom_column_clears_saved_guest_values(self):
        self.client.force_login(self.user)
        field = GuestField.objects.create(label="Dietary preference")
        self.guest.custom_data = {str(field.pk): "Vegetarian"}
        self.guest.save(update_fields=["custom_data"])
        self.client.post(reverse("delete_guest_field", args=[field.pk]))
        self.guest.refresh_from_db()
        self.assertFalse(GuestField.objects.filter(pk=field.pk).exists())
        self.assertNotIn(str(field.pk), self.guest.custom_data)

    def test_custom_column_limit_is_enforced(self):
        self.client.force_login(self.user)
        for index in range(10):
            GuestField.objects.create(label=f"Field {index}")
        self.client.post(reverse("add_guest_field"), {"label": "One too many"})
        self.assertEqual(GuestField.objects.count(), 10)

    def test_remove_demo_command_preserves_non_sample_records(self):
        self.client.force_login(self.user)
        call_command("seed_demo", stdout=io.StringIO())
        call_command("remove_demo", stdout=io.StringIO())
        self.assertEqual(list(Guest.objects.values_list("name", flat=True)), ["Riya Test"])

    @patch("guests.management.commands.create_organizer.getpass", side_effect=["Strong-passphrase-984!", "Strong-passphrase-984!"])
    @patch("builtins.input", return_value="new-organizer")
    def test_create_organizer_command_creates_staff_without_superuser(self, username_input, password_input):
        call_command("create_organizer", stdout=io.StringIO())
        user = get_user_model().objects.get(username="new-organizer")
        self.assertTrue(user.is_staff)
        self.assertFalse(user.is_superuser)
        self.assertTrue(user.check_password("Strong-passphrase-984!"))
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 200)

    @patch("guests.management.commands.create_organizer.getpass", side_effect=["Strong-passphrase-984!", "different-password"])
    @patch("builtins.input", return_value="uncreated-organizer")
    def test_create_organizer_command_rejects_mismatched_passwords(self, username_input, password_input):
        with self.assertRaises(CommandError):
            call_command("create_organizer", stdout=io.StringIO())
        self.assertFalse(
            get_user_model().objects.filter(username="uncreated-organizer").exists()
        )


class OrganizerSignupTests(TestCase):
    @override_settings(DEBUG=True, LOCAL_DEMO_MODE=False)
    def test_first_local_signup_creates_staff_and_signs_them_in(self):
        response = self.client.get(reverse("organizer_signup"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Create your account")

        response = self.client.post(
            reverse("organizer_signup"),
            {
                "username": "first-organizer",
                "password1": "Strong-passphrase-984!",
                "password2": "Strong-passphrase-984!",
            },
        )
        self.assertRedirects(response, reverse("dashboard"), fetch_redirect_response=False)
        user = get_user_model().objects.get(username="first-organizer")
        self.assertTrue(user.is_staff)
        self.assertTrue(user.check_password("Strong-passphrase-984!"))
        self.assertTrue(self.client.session.get("_auth_user_id"))

    @override_settings(DEBUG=True, LOCAL_DEMO_MODE=False)
    def test_signup_closes_after_first_staff_account(self):
        get_user_model().objects.create_user(
            username="existing-organizer",
            password="Strong-passphrase-984!",
            is_staff=True,
        )
        response = self.client.get(reverse("organizer_signup"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Account already set up")
        response = self.client.post(
            reverse("organizer_signup"),
            {
                "username": "second-organizer",
                "password1": "Strong-passphrase-984!",
                "password2": "Strong-passphrase-984!",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(
            get_user_model().objects.filter(username="second-organizer").exists()
        )

    @override_settings(DEBUG=True, LOCAL_DEMO_MODE=False)
    def test_signup_rejects_nonlocal_requests(self):
        response = self.client.get(
            reverse("organizer_signup"),
            REMOTE_ADDR="192.168.1.20",
        )
        self.assertEqual(response.status_code, 404)
