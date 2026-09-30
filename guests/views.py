import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from functools import wraps
from pathlib import Path
from urllib.parse import urlencode, urlparse
from zipfile import BadZipFile

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model, login
from django.contrib.auth.forms import UserCreationForm
from django.contrib.auth.views import redirect_to_login
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models import Count, Max, Q, Sum
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.core.paginator import Paginator
from django.shortcuts import get_object_or_404, redirect, render
from django.templatetags.static import static
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods, require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.defined_name import DefinedName
from PIL import Image, UnidentifiedImageError

from .country_codes import country_code_choices
from .models import (
    CrewMember,
    Event,
    EventMedia,
    Expense,
    Guest,
    GuestField,
    WhatsAppCampaign,
    WhatsAppCampaignRecipient,
)

MAX_EVENT_MEDIA_IMAGE_BYTES = 20 * 1024 * 1024
MAX_EVENT_MEDIA_VIDEO_BYTES = 100 * 1024 * 1024
MAX_EVENT_MEDIA_COUNT = 40
MAX_EVENT_MEDIA_TOTAL_BYTES = 500 * 1024 * 1024


def organizer_required(view_func):
    @wraps(view_func)
    def wrapped(request, *args, **kwargs):
        if (
            settings.LOCAL_DEMO_MODE
            and request.META.get("REMOTE_ADDR") in {"127.0.0.1", "::1"}
        ):
            return view_func(request, *args, **kwargs)
        if not request.user.is_authenticated:
            return redirect_to_login(request.get_full_path())
        if not request.user.is_staff:
            raise PermissionDenied
        return view_func(request, *args, **kwargs)

    return wrapped


def get_active_event(request):
    event_id = request.session.get("active_event_id")
    if event_id:
        event = Event.objects.filter(pk=event_id).first()
        if event:
            return event
    event = Event.objects.first()
    if event:
        request.session["active_event_id"] = event.pk
    return event


def workspace_context(request, **extra):
    return {
        "active_event": get_active_event(request),
        "events": Event.objects.all(),
        **extra,
    }


def whatsapp_invite_link(guest):
    base_url = settings.PUBLIC_BASE_URL
    parsed_url = urlparse(base_url)
    phone = guest_whatsapp_number(guest)
    if (
        not guest.whatsapp_opt_in
        or parsed_url.scheme != "https"
        or not parsed_url.netloc
        or not parsed_url.hostname
        or parsed_url.username
        or parsed_url.password
        or not phone
    ):
        return None

    rsvp_url = f"{base_url}{reverse('guest_rsvp', args=[guest.rsvp_token])}"
    message = (
        f"Hi {guest.name}, you're invited to {guest.event.name}. "
        f"Please RSVP here: {rsvp_url} "
        "Reply STOP if you no longer want event updates."
    )
    return f"https://wa.me/{phone}?{urlencode({'text': message})}"


def guest_whatsapp_number(guest):
    raw_phone = guest.phone.strip()
    phone_digits = re.sub(r"\D", "", raw_phone)
    country_digits = re.sub(r"\D", "", guest.country_code)
    if raw_phone.startswith("+"):
        if country_digits and not phone_digits.startswith(country_digits):
            return ""
    elif country_digits:
        phone_digits = country_digits + phone_digits
    else:
        return ""
    if re.fullmatch(r"[1-9]\d{7,14}", phone_digits):
        return phone_digits
    return ""


def normalize_contact_number(country_code, contact_number):
    country_code = str(country_code or "").strip()
    contact_number = str(contact_number or "").strip()
    if not re.fullmatch(r"\+?[1-9]\d{0,2}", country_code):
        return None
    country_digits = re.sub(r"\D", "", country_code)
    if not contact_number or not re.fullmatch(r"\+?[\d\s().-]+", contact_number):
        return None
    number_digits = re.sub(r"\D", "", contact_number)
    if contact_number.startswith("+"):
        if not number_digits.startswith(country_digits):
            return None
        number_digits = number_digits[len(country_digits):]
    international = country_digits + number_digits
    if not re.fullmatch(r"[1-9]\d{7,14}", international):
        return None
    return f"+{country_digits}", number_digits


def _clean_import_header(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def _excel_cell_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for format_string in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            if format_string == "%Y-%m-%d":
                return date.fromisoformat(text)
            return datetime.strptime(text, format_string).date()
        except ValueError:
            continue
    raise ValueError("use YYYY-MM-DD or DD/MM/YYYY")


@organizer_required
@require_http_methods(["GET", "POST"])
def event_list(request):
    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        venue = request.POST.get("venue", "").strip()
        start_value = request.POST.get("start_date", "").strip()
        end_value = request.POST.get("end_date", "").strip()
        if not name:
            messages.error(request, "Enter an event name.")
        elif len(name) > 150 or len(venue) > 200:
            messages.error(request, "Event name or venue is too long.")
        else:
            try:
                start_date = date.fromisoformat(start_value) if start_value else None
                end_date = date.fromisoformat(end_value) if end_value else None
            except ValueError:
                messages.error(request, "Enter valid event dates.")
            else:
                if start_date and end_date and end_date < start_date:
                    messages.error(request, "Event end date must be after its start date.")
                else:
                    event = Event.objects.create(
                        name=name,
                        venue=venue,
                        start_date=start_date,
                        end_date=end_date,
                    )
                    request.session["active_event_id"] = event.pk
                    messages.success(request, f"Created “{event.name}”.")
                    return redirect("event_list")
    return render(
        request,
        "guests/events.html",
        workspace_context(request, event_list=Event.objects.all()),
    )


@organizer_required
@require_POST
def select_event(request, event_id):
    event = get_object_or_404(Event, pk=event_id)
    request.session["active_event_id"] = event.pk
    messages.success(request, f"Now managing “{event.name}”.")
    return redirect("dashboard")


@organizer_required
@require_http_methods(["GET", "POST"])
def expense_list(request):
    event = get_active_event(request)
    if not event:
        messages.error(request, "Create an event before adding expenses.")
        return redirect("event_list")
    if request.method == "POST":
        title = request.POST.get("title", "").strip()
        category = request.POST.get("category", "").strip()
        amount_value = request.POST.get("amount", "").strip()
        spent_on_value = request.POST.get("spent_on", "").strip()
        notes = request.POST.get("notes", "").strip()
        try:
            amount = Decimal(amount_value)
            spent_on = date.fromisoformat(spent_on_value) if spent_on_value else None
        except (InvalidOperation, ValueError):
            messages.error(request, "Enter a valid expense amount and date.")
        else:
            if not title or len(title) > 150 or len(category) > 80 or len(notes) > 500:
                messages.error(request, "Enter a valid title, category, and notes.")
            elif (
                not amount.is_finite()
                or amount <= 0
                or amount > Decimal("9999999999.99")
                or amount.quantize(Decimal("0.01")) != amount
            ):
                messages.error(request, "Expense amount must be positive with up to 2 decimals.")
            else:
                Expense.objects.create(
                    event=event,
                    title=title,
                    category=category,
                    amount=amount,
                    spent_on=spent_on,
                    notes=notes,
                )
                messages.success(request, f"Added expense “{title}”.")
                return redirect("expense_list")

    expenses = event.expenses.all()
    return render(
        request,
        "guests/expenses.html",
        workspace_context(
            request,
            expenses=expenses,
            expense_total=expenses.aggregate(total=Sum("amount"))["total"] or Decimal("0.00"),
        ),
    )


@organizer_required
@require_POST
def delete_expense(request, expense_id):
    expense = get_object_or_404(
        Expense,
        pk=expense_id,
        event=get_active_event(request),
    )
    title = expense.title
    expense.delete()
    messages.success(request, f"Removed expense “{title}”.")
    return redirect("expense_list")


@organizer_required
@require_http_methods(["GET", "POST"])
def crew_list(request):
    event = get_active_event(request)
    if not event:
        messages.error(request, "Create an event before adding crew members.")
        return redirect("event_list")
    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        role = request.POST.get("role", "").strip()
        phone = request.POST.get("phone", "").strip()
        email = request.POST.get("email", "").strip()
        if not name or not role:
            messages.error(request, "Enter the crew member’s name and role.")
        elif len(name) > 150 or len(role) > 100 or len(phone) > 32:
            messages.error(request, "One or more crew details are too long.")
        else:
            member = CrewMember(
                event=event,
                name=name,
                role=role,
                phone=phone,
                email=email,
            )
            try:
                member.full_clean()
            except ValidationError:
                messages.error(request, "Enter a valid email address.")
            else:
                member.save()
                messages.success(request, f"Added {member.name} to the crew.")
                return redirect("crew_list")

    return render(
        request,
        "guests/crew.html",
        workspace_context(request, crew_members=event.crew.all()),
    )


@organizer_required
@require_POST
def delete_crew_member(request, member_id):
    member = get_object_or_404(
        CrewMember,
        pk=member_id,
        event=get_active_event(request),
    )
    name = member.name
    member.delete()
    messages.success(request, f"Removed {name} from the crew.")
    return redirect("crew_list")


@organizer_required
def dashboard(request):
    active_event = get_active_event(request)
    event_guests = (
        Guest.objects.filter(event=active_event)
        if active_event
        else Guest.objects.none()
    )
    counts = event_guests.aggregate(
        total=Count("id"),
        confirmed=Count("id", filter=Q(rsvp_status=Guest.RSVPStatus.CONFIRMED)),
        pending=Count("id", filter=Q(rsvp_status=Guest.RSVPStatus.PENDING)),
        rooms=Count(
            "id",
            filter=~Q(hotel_status=Guest.HotelStatus.NOT_SELECTED)
            & ~Q(rsvp_status=Guest.RSVPStatus.DECLINED),
        ),
    )
    side_counts = event_guests.aggregate(
        bride=Count("id", filter=Q(side_tag=Guest.SideTag.BRIDE)),
        groom=Count("id", filter=Q(side_tag=Guest.SideTag.GROOM)),
        unassigned=Count("id", filter=Q(side_tag="")),
    )
    guests = event_guests
    query = request.GET.get("q", "").strip()
    status = request.GET.get("status", "")
    side = request.GET.get("side", "")
    if side == Guest.SideTag.BRIDE:
        guests = guests.filter(side_tag=Guest.SideTag.BRIDE)
    elif side == Guest.SideTag.GROOM:
        guests = guests.filter(side_tag=Guest.SideTag.GROOM)
    elif side == "unassigned":
        guests = guests.filter(side_tag="")
    else:
        side = ""
    if query:
        guests = guests.filter(
            Q(name__icontains=query)
            | Q(last_name__icontains=query)
            | Q(phone__icontains=query)
        )
    if status in Guest.RSVPStatus.values:
        guests = guests.filter(rsvp_status=status)
    guests = guests.order_by("side_tag", "group_serial", "name", "last_name", "id")
    event_media = active_event.media.all()[:8] if active_event else []
    event_media_count = active_event.media.count() if active_event else 0

    custom_fields = list(GuestField.objects.all())
    guest_page = Paginator(guests, 100).get_page(request.GET.get("page"))
    for guest in guest_page.object_list:
        guest.custom_values = [
            {
                "field": field,
                "value": guest.custom_data.get(str(field.pk), ""),
            }
            for field in custom_fields
        ]
        guest.whatsapp_invite_link = whatsapp_invite_link(guest)
        guest.whatsapp_number_valid = bool(guest_whatsapp_number(guest))

    context = workspace_context(
        request,
        guests=guest_page.object_list,
        event_media=event_media,
        event_media_count=event_media_count,
        guest_page=guest_page,
        custom_fields=custom_fields,
        column_count=14 + len(custom_fields),
        custom_field_limit=10,
        counts=counts,
        query=query,
        selected_status=status,
        status_choices=Guest.RSVPStatus.choices,
        hotel_choices=Guest.HotelStatus.choices,
        room_choices=Guest.RoomPreference.choices,
        guest_type_choices=Guest.GuestType.choices,
        side_tag_choices=Guest.SideTag.choices,
        country_code_choices=country_code_choices(),
        sub_category_choices=Guest.SubCategory.choices,
        side_counts=side_counts,
        selected_side=side,
        whatsapp_business_configured=all(
            (
                settings.WHATSAPP_PHONE_NUMBER_ID,
                settings.WHATSAPP_ACCESS_TOKEN,
                settings.WHATSAPP_API_VERSION,
            )
        ),
        whatsapp_campaigns=active_event.whatsapp_campaigns.all()[:5]
        if active_event
        else [],
        public_rsvp_enabled=bool(settings.PUBLIC_BASE_URL),
        event_expense_total=active_event.expenses.aggregate(total=Sum("amount"))["total"]
        if active_event
        else 0,
        signup_available=not get_user_model().objects.filter(is_staff=True).exists(),
        confirmed_percent=round(counts["confirmed"] * 100 / counts["total"])
        if counts["total"]
        else 0,
    )
    return render(request, "guests/dashboard.html", context)


def _validate_event_media(upload):
    extension = Path(upload.name).suffix.lower()
    image_types = {
        ".jpg": ("JPEG", "image/jpeg"),
        ".jpeg": ("JPEG", "image/jpeg"),
        ".png": ("PNG", "image/png"),
        ".webp": ("WEBP", "image/webp"),
    }
    video_types = {
        ".mp4": "video/mp4",
        ".webm": "video/webm",
    }
    if extension in image_types:
        if upload.size > MAX_EVENT_MEDIA_IMAGE_BYTES:
            raise ValueError("Photos must be 20 MB or smaller.")
        expected_format, content_type = image_types[extension]
        try:
            with Image.open(upload) as image:
                if image.format != expected_format:
                    raise ValueError("The photo content does not match its file type.")
                if image.width * image.height > 40_000_000:
                    raise ValueError("The photo dimensions are too large.")
                image.verify()
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as error:
            raise ValueError("The photo file is damaged or not a supported image.") from error
        upload.seek(0)
        return EventMedia.MediaType.IMAGE, content_type
    if extension in video_types:
        if upload.size > MAX_EVENT_MEDIA_VIDEO_BYTES:
            raise ValueError("Videos must be 100 MB or smaller.")
        signature = upload.read(16)
        upload.seek(0)
        valid_video = (
            len(signature) >= 12 and signature[4:8] == b"ftyp"
            if extension == ".mp4"
            else signature.startswith(b"\x1a\x45\xdf\xa3")
        )
        if not valid_video:
            raise ValueError("The video file content does not match MP4 or WebM.")
        return EventMedia.MediaType.VIDEO, video_types[extension]
    raise ValueError("Choose a JPG, PNG, or WebP photo, or an MP4/WebM video.")


@organizer_required
@require_POST
def upload_event_media(request):
    event = get_active_event(request)
    if not event:
        messages.error(request, "Create or select an event before adding media.")
        return redirect("event_list")
    upload = request.FILES.get("media_file")
    if upload is None:
        messages.error(request, "Choose a photo or video to upload.")
        return redirect("dashboard")
    if event.media.count() >= MAX_EVENT_MEDIA_COUNT:
        messages.error(request, f"An event can have up to {MAX_EVENT_MEDIA_COUNT} media items.")
        return redirect("dashboard")
    stored_bytes = sum(item.file.size for item in event.media.all())
    if stored_bytes + upload.size > MAX_EVENT_MEDIA_TOTAL_BYTES:
        messages.error(request, "Event media storage is limited to 500 MB.")
        return redirect("dashboard")
    try:
        media_type, content_type = _validate_event_media(upload)
    except ValueError as error:
        messages.error(request, str(error))
        return redirect("dashboard")
    title = request.POST.get("title", "").strip()
    if len(title) > 100:
        messages.error(request, "Media captions must be 100 characters or fewer.")
        return redirect("dashboard")
    media = EventMedia(
        event=event,
        title=title,
        media_type=media_type,
        content_type=content_type,
        uploaded_by=request.user if request.user.is_authenticated else None,
    )
    media.file.save(upload.name, upload, save=False)
    media.save()
    messages.success(request, "Event media uploaded.")
    return redirect("dashboard")


@organizer_required
@require_POST
def delete_event_media(request, media_id):
    event = get_active_event(request)
    if not event:
        messages.error(request, "Select an event before removing media.")
        return redirect("event_list")
    media = get_object_or_404(EventMedia, pk=media_id, event=event)
    title = media.title or media.get_media_type_display()
    media.delete()
    messages.success(request, f"Removed {title} from this event.")
    return redirect("dashboard")


@organizer_required
def event_media_file(request, media_id):
    media = get_object_or_404(
        EventMedia,
        pk=media_id,
        event=get_active_event(request),
    )
    if media.content_type not in {
        "image/jpeg",
        "image/png",
        "image/webp",
        "video/mp4",
        "video/webm",
    }:
        raise Http404
    try:
        response = FileResponse(
            media.file.open("rb"),
            content_type=media.content_type,
        )
    except OSError as error:
        raise Http404 from error
    response["Cache-Control"] = "private, no-store"
    response["X-Content-Type-Options"] = "nosniff"
    response["Content-Security-Policy"] = "default-src 'none'; sandbox"
    return response


@organizer_required
@require_POST
def add_guest(request):
    name = request.POST.get("name", "").strip()
    last_name = request.POST.get("last_name", "").strip()
    phone = request.POST.get("phone", "").strip()
    country_code_input = request.POST.get("country_code", "").strip()
    guest_type = request.POST.get("guest_type", "")
    side_tag = request.POST.get("side_tag", "")
    group_serial_input = request.POST.get("group_serial", "").strip()
    consent = request.POST.get("message_consent") == "yes"
    active_event = get_active_event(request)
    if not active_event:
        messages.error(request, "Create an event before adding guests.")
    elif guest_type not in {"", *Guest.GuestType.values}:
        messages.error(request, "Choose a valid guest type.")
    elif side_tag not in {"", *Guest.SideTag.values}:
        messages.error(request, "Choose Bride’s side or Groom’s side, or leave it unassigned.")
    elif (
        len(name) > 150
        or len(last_name) > 150
        or len(phone) > 32
        or len(country_code_input) > 6
    ):
        messages.error(request, "The name or contact details are too long.")
    else:
        normalized_contact = (
            normalize_contact_number(country_code_input, phone)
            if phone and country_code_input
            else None
        )
        try:
            group_serial = (
                int(group_serial_input)
                if group_serial_input
                else (
                    Guest.objects.filter(
                        event=active_event,
                        side_tag=side_tag,
                    ).aggregate(max_serial=Max("group_serial"))["max_serial"]
                    or 0
                )
                + 1
            )
        except ValueError:
            group_serial = 0
        if group_serial < 1:
            messages.error(request, "Group ID must be a positive whole number.")
        else:
            if normalized_contact:
                country_code, phone = normalized_contact
            else:
                country_code = country_code_input
            Guest.objects.create(
                event=active_event,
                group_serial=group_serial,
                name=name,
                last_name=last_name,
                phone=phone,
                country_code=country_code,
                guest_type=guest_type,
                side_tag=side_tag,
                whatsapp_opt_in=bool(normalized_contact and consent),
                whatsapp_opt_in_at=(
                    timezone.now() if normalized_contact and consent else None
                ),
            )
            display_name = " ".join(part for part in (name, last_name) if part) or "Guest"
            messages.success(request, f"Added {display_name} to Group ID {group_serial}.")
    return redirect("dashboard")


@organizer_required
@require_POST
def add_guest_field(request):
    label = request.POST.get("label", "").strip()
    if not label:
        messages.error(request, "Enter a name for the new guest-list column.")
    elif len(label) > 50:
        messages.error(request, "Column names must be 50 characters or fewer.")
    elif GuestField.objects.filter(label__iexact=label).exists():
        messages.error(request, "A column with that name already exists.")
    elif GuestField.objects.count() >= 10:
        messages.error(request, "You can add up to 10 custom guest-list columns.")
    else:
        GuestField.objects.create(label=label)
        messages.success(
            request,
            f"Added the “{label}” column. Existing guest values start blank.",
        )
    return redirect("dashboard")


@organizer_required
@require_POST
def delete_guest_field(request, field_id):
    field = get_object_or_404(GuestField, pk=field_id)
    with transaction.atomic():
        for guest in Guest.objects.only("id", "custom_data").iterator():
            if str(field.pk) in guest.custom_data:
                guest.custom_data.pop(str(field.pk), None)
                guest.save(update_fields=["custom_data"])
        label = field.label
        field.delete()
    messages.success(request, f"Removed “{label}” and its guest values.")
    return redirect("dashboard")


@organizer_required
@require_POST
def import_guests(request):
    active_event = get_active_event(request)
    if not active_event:
        messages.error(request, "Create an event before importing guests.")
        return redirect("event_list")
    upload = request.FILES.get("guest_file")
    if upload is None:
        messages.error(request, "Choose an Excel (.xlsx) or CSV file to import.")
        return redirect("dashboard")
    if upload.size > 2 * 1024 * 1024:
        messages.error(request, "The import file must be 2 MB or smaller.")
        return redirect("dashboard")
    extension = Path(upload.name).suffix.lower()
    requested_default_side = request.POST.get("default_side", "").strip()
    import_default_side = {
        Guest.SideTag.BRIDE: Guest.SideTag.BRIDE,
        Guest.SideTag.GROOM: Guest.SideTag.GROOM,
    }.get(requested_default_side)
    sheet_headers = {}
    try:
        if extension == ".xlsx":
            workbook = load_workbook(upload, read_only=True, data_only=True)
            sheet_rows = []
            side_sheet_names = {
                "bride side": Guest.SideTag.BRIDE,
                "brides side": Guest.SideTag.BRIDE,
                "bride": Guest.SideTag.BRIDE,
                "groom side": Guest.SideTag.GROOM,
                "grooms side": Guest.SideTag.GROOM,
                "groom": Guest.SideTag.GROOM,
            }
            ignored_sheet_names = {"instructions", "read me", "readme", "lists"}
            for sheet in workbook.worksheets:
                normalized_sheet_name = _clean_import_header(sheet.title)
                if normalized_sheet_name in ignored_sheet_names:
                    continue
                default_side = side_sheet_names.get(normalized_sheet_name)
                all_rows = sheet.iter_rows(values_only=True)
                raw_headers = next(all_rows, ())
                if not raw_headers or not any(
                    _clean_import_header(value) for value in raw_headers
                ):
                    continue
                if len(raw_headers) > 100:
                    messages.error(request, "A worksheet cannot contain more than 100 columns.")
                    return redirect("dashboard")
                header_names = [_clean_import_header(value) for value in raw_headers]
                if len(set(header for header in header_names if header)) != len(
                    [header for header in header_names if header]
                ):
                    messages.error(request, f"The {sheet.title} sheet has duplicate column names.")
                    return redirect("dashboard")
                sheet_headers[sheet.title] = {
                    header: header for header in header_names if header
                }
                for line_number, values in enumerate(all_rows, start=2):
                    if any(value not in (None, "") for value in values):
                        if len(sheet_rows) >= 2000:
                            messages.error(request, "A single import cannot exceed 2,000 guest rows.")
                            return redirect("dashboard")
                        sheet_rows.append(
                            (
                                default_side,
                                sheet.title,
                                line_number,
                                dict(zip(header_names, values)),
                            )
                        )
            if not sheet_headers:
                messages.error(
                    request,
                    "The workbook has no guest-list sheets. Include guest rows in Bride Side, Groom Side, or a sheet with a Tag/Side column.",
                )
                return redirect("dashboard")
        elif extension == ".csv":
            text = upload.read().decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text))
            raw_headers = next(reader, ())
            if len(raw_headers) > 100:
                messages.error(request, "The CSV cannot contain more than 100 columns.")
                return redirect("dashboard")
            header_names = [_clean_import_header(value) for value in raw_headers]
            if len(set(header for header in header_names if header)) != len(
                [header for header in header_names if header]
            ):
                messages.error(request, "The CSV has duplicate column names.")
                return redirect("dashboard")
            data_rows = []
            for line_number, values in enumerate(reader, start=2):
                if any(value not in (None, "") for value in values):
                    data_rows.append(
                        (None, "CSV", line_number, dict(zip(header_names, values)))
                    )
        else:
            messages.error(request, "Choose an Excel (.xlsx) or CSV file.")
            return redirect("dashboard")
    except (BadZipFile, UnicodeDecodeError, InvalidFileException, OSError, ValueError):
        messages.error(request, "The file could not be read. Check it and try again.")
        return redirect("dashboard")
    finally:
        if extension == ".xlsx" and "workbook" in locals():
            workbook.close()

    if extension == ".xlsx":
        data_rows = sheet_rows
    else:
        sheet_headers = {
            "CSV": {
                header: header
                for header in header_names
                if header
            }
        }

    field_aliases = {
        "group_serial": (
            "group id", "serial number", "serial no", "serial", "group serial", "s no",
            "serial number group", "family serial",
        ),
        "name": ("first name", "name", "guest name", "guest first name"),
        "last_name": ("last name", "surname", "family name"),
        "country_code": (
            "country code", "country calling code", "dial code", "country dial code",
        ),
        "phone": (
            "contact number", "whatsapp number", "whatsapp contact", "phone",
            "mobile", "phone number", "contact",
        ),
        "guest_type": ("guest type", "type"),
        "invited_date": ("invited date", "invited dates", "invite date", "date invited"),
        "category": ("category", "relation", "relationship", "category relation"),
        "pax": ("pax", "guests", "guest count", "number of guests"),
        "sub_category": ("sub category", "subcategory"),
        "invite_name": ("invite name", "invitation name"),
        "salutation": ("salutation", "title"),
        "side_tag": (
            "tag", "side", "bride groom", "bride or groom", "guest side",
            "bride groom side", "bride groom tag", "family side",
            "bride groom tag side",
        ),
    }
    known_header_names = {
        alias for aliases in field_aliases.values() for alias in aliases
    }
    all_custom_headers = {
        header.title()
        for sheet_header_map in sheet_headers.values()
        for header in sheet_header_map
        if header and header not in known_header_names
    }
    guests_to_create = []
    errors = []
    if len(data_rows) > 2000:
        messages.error(request, "A single import cannot exceed 2,000 guest rows.")
        return redirect("dashboard")
    side_value_map = {
        "bride": Guest.SideTag.BRIDE,
        "brides side": Guest.SideTag.BRIDE,
        "brides": Guest.SideTag.BRIDE,
        "bride side": Guest.SideTag.BRIDE,
        "bride s side": Guest.SideTag.BRIDE,
        "groom": Guest.SideTag.GROOM,
        "grooms side": Guest.SideTag.GROOM,
        "grooms": Guest.SideTag.GROOM,
        "groom side": Guest.SideTag.GROOM,
        "groom s side": Guest.SideTag.GROOM,
    }
    next_group_id = dict(
        Guest.objects.filter(event=active_event)
        .values("side_tag")
        .annotate(max_group_id=Max("group_serial"))
        .values_list("side_tag", "max_group_id")
    )
    imported_group_max = {}
    for sheet_side, sheet_title, _, row in data_rows:
        normalized_headers = sheet_headers[sheet_title]
        group_column = next(
            (
                normalized_headers[alias]
                for alias in field_aliases["group_serial"]
                if alias in normalized_headers
            ),
            None,
        )
        side_column = next(
            (
                normalized_headers[alias]
                for alias in field_aliases["side_tag"]
                if alias in normalized_headers
            ),
            None,
        )
        side_raw = str(row.get(side_column) or "").strip() if side_column else ""
        group_side = (
            side_value_map.get(_clean_import_header(side_raw), "")
            if side_raw
            else (sheet_side or import_default_side or "")
        )
        group_value = row.get(group_column) if group_column else None
        try:
            group_number = int(float(group_value))
            if float(group_value) == group_number and group_number > 0:
                imported_group_max[group_side] = max(
                    imported_group_max.get(group_side, 0),
                    group_number,
                )
        except (TypeError, ValueError, OverflowError):
            continue
    for group_side, group_number in imported_group_max.items():
        next_group_id[group_side] = max(next_group_id.get(group_side) or 0, group_number)
    for sheet_side, sheet_title, line_number, row in data_rows:
        normalized_headers = sheet_headers[sheet_title]
        columns = {
            field: next(
                (
                    normalized_headers[alias]
                    for alias in aliases
                    if alias in normalized_headers
                ),
                None,
            )
            for field, aliases in field_aliases.items()
        }
        name = str(row.get(columns["name"]) or "").strip()
        last_name_column = columns["last_name"]
        last_name = str(row.get(last_name_column) or "").strip() if last_name_column else ""
        if len(name) > 150 or len(last_name) > 150:
            errors.append(f"{sheet_title} row {line_number}: First or Last Name exceeds 150 characters.")
            continue
        side_raw = str(row.get(columns["side_tag"]) or "").strip()
        if side_raw:
            side_tag = side_value_map.get(_clean_import_header(side_raw), "")
        else:
            side_tag = sheet_side or import_default_side or ""

        serial_value = row.get(columns["group_serial"])
        if serial_value in (None, ""):
            group_serial = (next_group_id.get(side_tag) or 0) + 1
            next_group_id[side_tag] = group_serial
        else:
            try:
                serial_numeric = float(serial_value)
                group_serial = int(serial_numeric)
                if serial_numeric != group_serial or group_serial < 1:
                    raise ValueError
            except (TypeError, ValueError, OverflowError):
                group_serial = (next_group_id.get(side_tag) or 0) + 1
                next_group_id[side_tag] = group_serial
                invalid_group_value = str(serial_value).strip()
            else:
                invalid_group_value = ""
        if serial_value in (None, ""):
            invalid_group_value = ""

        country_raw = str(row.get(columns["country_code"]) or "").strip()
        if len(country_raw) > Guest._meta.get_field("country_code").max_length:
            errors.append(f"{sheet_title} row {line_number}: Country Code is too long.")
            continue
        country_digits = re.sub(r"\D", "", country_raw)
        country_code = (
            f"+{country_digits}"
            if re.fullmatch(r"\+?[1-9]\d{0,2}", country_raw)
            else country_raw
        )
        phone_raw = str(row.get(columns["phone"]) or "").strip()
        if len(phone_raw) > Guest._meta.get_field("phone").max_length:
            errors.append(f"{sheet_title} row {line_number}: Phone Number is too long.")
            continue
        phone = phone_raw
        if phone_raw and country_digits and re.fullmatch(r"\+?[\d\s().-]+", phone_raw):
            phone_digits = re.sub(r"\D", "", phone_raw)
            if phone_raw.startswith("+"):
                phone_digits = (
                    phone_digits[len(country_digits):]
                    if phone_digits.startswith(country_digits)
                    else ""
                )
            if phone_digits and re.fullmatch(r"[1-9]\d{7,14}", country_digits + phone_digits):
                phone = phone_digits

        guest_type_raw = str(row.get(columns["guest_type"]) or "").strip()
        guest_type_map = {
            _clean_import_header(label): key
            for key, label in Guest.GuestType.choices
        }
        guest_type_map.update({key: key for key, _ in Guest.GuestType.choices})
        guest_type_map.update({
            "local": Guest.GuestType.LOCAL_NON_STAY,
            "local non stay": Guest.GuestType.LOCAL_NON_STAY,
            "local stay": Guest.GuestType.LOCAL_STAY,
        })
        guest_type = guest_type_map.get(_clean_import_header(guest_type_raw), "")
        fields = {
            "event": active_event,
            "group_serial": group_serial,
            "name": name,
            "phone": phone,
            "country_code": country_code,
            "guest_type": guest_type,
            "side_tag": side_tag,
        }
        fields["last_name"] = last_name
        if guest_type_raw and not guest_type:
            fields.setdefault("_custom_values", {})["Provided Guest Type"] = guest_type_raw
        if invalid_group_value:
            fields.setdefault("_custom_values", {})["Provided Group ID"] = invalid_group_value
        if side_raw and not side_tag:
            fields.setdefault("_custom_values", {})["Provided Tag"] = side_raw
        row_error = None
        for field in ("invited_date", "category", "pax", "sub_category", "invite_name", "salutation"):
            column = columns[field]
            if not column:
                continue
            raw_value = row.get(column)
            if raw_value in (None, ""):
                continue
            if field == "invited_date":
                try:
                    fields[field] = _excel_cell_date(raw_value)
                except ValueError:
                    fields.setdefault("_custom_values", {})["Provided Invited Dates"] = str(raw_value).strip()
                continue
            if field == "pax":
                if not str(raw_value).strip():
                    fields[field] = None
                    continue
                try:
                    pax_numeric = float(raw_value)
                    pax = int(pax_numeric)
                    if pax_numeric != pax or not 1 <= pax <= 100:
                        raise ValueError
                    fields[field] = pax
                except (TypeError, ValueError, OverflowError):
                    fields.setdefault("_custom_values", {})["Provided Pax"] = str(raw_value).strip()
                continue
            value = str(raw_value).strip()
            if field == "sub_category":
                sub_category_map = {
                    "single": "single",
                    "couple": "couple",
                }
                value = sub_category_map.get(_clean_import_header(value), "")
                if not value:
                    fields.setdefault("_custom_values", {})["Provided Sub Category"] = str(raw_value).strip()
                    continue
            if len(value) > Guest._meta.get_field(field).max_length:
                row_error = f"{sheet_title} row {line_number}: {field.replace('_', ' ').title()} is too long."
                break
            fields[field] = value
        if row_error:
            errors.append(row_error)
            continue
        for header, column in normalized_headers.items():
            if header in known_header_names:
                continue
            raw_value = row.get(column)
            value = "" if raw_value is None else str(raw_value).strip()
            if not value:
                continue
            label = column.strip().title()
            if len(label) > 50 or len(value) > 500:
                row_error = f"{sheet_title} row {line_number}: extra detail field or value is too long."
                break
            fields.setdefault("_custom_values", {})[label] = value
        if row_error:
            errors.append(row_error)
            continue
        guests_to_create.append((Guest(**{key: value for key, value in fields.items() if key != "_custom_values"}), fields.get("_custom_values", {}), (side_tag, group_serial), sheet_title, line_number))

    if errors:
        messages.error(request, "Import stopped: " + " ".join(errors[:5]))
        return redirect("dashboard")
    if not guests_to_create:
        messages.error(request, "The workbook has no guest rows to import.")
        return redirect("dashboard")

    available_custom_fields = {field.label.lower(): field for field in GuestField.objects.all()}
    new_labels = {
        label for label in all_custom_headers if label.lower() not in available_custom_fields
    }
    new_labels.update(
        label
        for _, custom_values, _, _, _ in guests_to_create
        for label in custom_values
        if label.lower() not in available_custom_fields
    )
    if len(available_custom_fields) + len(new_labels) > 10:
        messages.error(request, "Import has more extra columns than the 10-column limit.")
        return redirect("dashboard")
    with transaction.atomic():
        for label in sorted(new_labels):
            available_custom_fields[label.lower()] = GuestField.objects.create(label=label)
        new_guests = []
        for guest, custom_values, _, _, _ in guests_to_create:
            guest.custom_data = {
                str(available_custom_fields[label.lower()].pk): value
                for label, value in custom_values.items()
            }
            new_guests.append(guest)
        Guest.objects.bulk_create(new_guests)
    messages.success(
        request,
        f"Imported {len(new_guests)} guests. Missing Group IDs were assigned automatically. "
        "Imported contacts default to no WhatsApp consent.",
    )
    return redirect("dashboard")


@organizer_required
@require_POST
def update_guest(request, guest_id):
    guest = get_object_or_404(
        Guest,
        pk=guest_id,
        event=get_active_event(request),
    )
    status = request.POST.get("rsvp_status", guest.rsvp_status)
    hotel_status = request.POST.get("hotel_status", guest.hotel_status)
    opt_in_value = request.POST.get(
        "whatsapp_opt_in", "yes" if guest.whatsapp_opt_in else "no"
    )
    name = request.POST.get("name", guest.name).strip()
    last_name = request.POST.get("last_name", guest.last_name).strip()
    phone_input = request.POST.get("phone", guest.phone).strip()
    country_code_input = request.POST.get("country_code", guest.country_code).strip()
    side_tag = request.POST.get("side_tag", guest.side_tag)
    group_serial_input = request.POST.get(
        "group_serial",
        str(guest.group_serial) if guest.group_serial is not None else "",
    ).strip()
    if (
        status not in Guest.RSVPStatus.values
        or hotel_status not in Guest.HotelStatus.values
        or opt_in_value not in {"yes", "no"}
    ):
        messages.error(request, "Choose a valid RSVP and hotel status.")
    else:
        try:
            group_serial = int(group_serial_input) if group_serial_input else None
        except (ValueError, TypeError):
            group_serial = -1
        normalized_contact = (
            normalize_contact_number(country_code_input, phone_input)
            if phone_input and country_code_input
            else None
        )
        if len(name) > 150 or len(last_name) > 150 or group_serial == -1 or (
            group_serial is not None and group_serial < 1
        ):
            messages.error(request, "Names must be 150 characters or fewer and Group ID must be positive when provided.")
            return redirect("dashboard")
        if side_tag not in {"", *Guest.SideTag.values}:
            messages.error(request, "Choose Bride Side or Groom Side.")
            return redirect("dashboard")
        if len(phone_input) > 32:
            messages.error(request, "Phone Number must be 32 characters or fewer.")
            return redirect("dashboard")
        guest.name = name
        guest.last_name = last_name
        guest.group_serial = group_serial
        guest.side_tag = side_tag
        if normalized_contact:
            guest.country_code, guest.phone = normalized_contact
        else:
            guest.phone = phone_input
        invited_date_value = request.POST.get("invited_date", "").strip()
        pax_value = request.POST.get("pax", str(guest.pax or "")).strip()
        guest_type = request.POST.get("guest_type", guest.guest_type)
        sub_category = request.POST.get("sub_category", "")
        try:
            invited_date = date.fromisoformat(invited_date_value) if invited_date_value else None
            pax = int(pax_value) if pax_value else None
        except (ValueError, TypeError):
            messages.error(request, "Enter a valid invited date and pax count.")
            return redirect("dashboard")
        if (
            (pax is not None and not 1 <= pax <= 100)
            or guest_type not in {"", *Guest.GuestType.values}
            or sub_category not in {"", *Guest.SubCategory.values}
        ):
            messages.error(request, "Choose valid guest details and pax from 1 to 100.")
            return redirect("dashboard")
        text_fields = {
            "country_code": request.POST.get("country_code", "").strip(),
            "category": request.POST.get("category", "").strip(),
            "invite_name": request.POST.get("invite_name", "").strip(),
            "salutation": request.POST.get("salutation", "").strip(),
        }
        for field, value in text_fields.items():
            if len(value) > Guest._meta.get_field(field).max_length:
                messages.error(request, f"{field.replace('_', ' ').title()} is too long.")
                return redirect("dashboard")
        custom_data = {}
        for field in GuestField.objects.all():
            value = request.POST.get(f"custom_{field.pk}", "").strip()
            if len(value) > 500:
                messages.error(
                    request, f"“{field.label}” must be 500 characters or fewer."
                )
                return redirect("dashboard")
            if value:
                custom_data[str(field.pk)] = value
        guest.rsvp_status = status
        guest.hotel_status = hotel_status
        guest.country_code = text_fields["country_code"]
        guest.invited_date = invited_date
        guest.category = text_fields["category"]
        guest.pax = pax
        guest.guest_type = guest_type
        guest.sub_category = sub_category
        guest.invite_name = text_fields["invite_name"]
        guest.salutation = text_fields["salutation"]
        guest.custom_data = custom_data
        if opt_in_value == "yes" and not guest.whatsapp_opt_in:
            guest.whatsapp_opt_in_at = timezone.now()
            guest.whatsapp_opted_out_at = None
        elif opt_in_value == "no" and guest.whatsapp_opt_in:
            guest.whatsapp_opted_out_at = timezone.now()
        guest.whatsapp_opt_in = opt_in_value == "yes"
        guest.save()
        guest_label = " ".join(part for part in (guest.name, guest.last_name) if part) or "Guest"
        messages.success(request, f"Updated {guest_label}'s details.")
    return redirect("dashboard")


@organizer_required
@require_POST
def queue_whatsapp_campaign(request):
    event = get_active_event(request)
    if not event:
        messages.error(request, "Create an event before queueing WhatsApp messages.")
        return redirect("event_list")
    if not all(
        (
            settings.WHATSAPP_PHONE_NUMBER_ID,
            settings.WHATSAPP_ACCESS_TOKEN,
            settings.WHATSAPP_API_VERSION,
        )
    ):
        messages.error(request, "Configure the WhatsApp Business API credentials first.")
        return redirect("dashboard")
    if (
        not re.fullmatch(r"\d+", settings.WHATSAPP_PHONE_NUMBER_ID)
        or not re.fullmatch(r"v\d+\.\d+", settings.WHATSAPP_API_VERSION)
    ):
        messages.error(request, "WhatsApp Business API configuration is invalid.")
        return redirect("dashboard")
    if urlparse(settings.PUBLIC_BASE_URL).scheme != "https":
        messages.error(request, "A public HTTPS RSVP URL is required before messaging.")
        return redirect("dashboard")

    template_name = request.POST.get("template_name", "").strip()
    language_code = request.POST.get("language_code", "").strip()
    if not re.fullmatch(r"[a-z0-9_]{1,512}", template_name):
        messages.error(request, "Enter the exact name of an approved WhatsApp template.")
        return redirect("dashboard")
    if not re.fullmatch(r"[A-Za-z0-9_-]{2,16}", language_code):
        messages.error(request, "Enter a valid WhatsApp template language code.")
        return redirect("dashboard")

    eligible_guests = [
        guest
        for guest in Guest.objects.filter(
            event=event,
            whatsapp_opt_in=True,
            whatsapp_opted_out_at__isnull=True,
        ).order_by("id")
        if guest_whatsapp_number(guest)
    ]
    if not eligible_guests:
        messages.error(
            request,
            "No opted-in guests with valid international phone numbers are eligible.",
        )
        return redirect("dashboard")
    if len(eligible_guests) > 5000:
        messages.error(request, "A single campaign cannot exceed 5,000 recipients.")
        return redirect("dashboard")

    with transaction.atomic():
        campaign = WhatsAppCampaign.objects.create(
            event=event,
            template_name=template_name,
            language_code=language_code,
            recipient_count=len(eligible_guests),
        )
        WhatsAppCampaignRecipient.objects.bulk_create(
            [
                WhatsAppCampaignRecipient(campaign=campaign, guest=guest)
                for guest in eligible_guests
            ]
        )
    messages.success(
        request,
        f"Queued {len(eligible_guests)} individual WhatsApp messages. "
        "A server operator must run the sender command to deliver them.",
    )
    return redirect("dashboard")


def _safe_csv_cell(value):
    value = str(value)
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _guest_export_headers(custom_fields):
    return [
        "Group ID",
        "First Name",
        "Last Name",
        "Country Code",
        "Phone Number",
        "Invited Dates",
        "Category",
        "Pax",
        "Guest Type",
        "Tag",
        "Sub Category",
        "Invite Name",
        "Salutation",
        *[_safe_csv_cell(field.label) for field in custom_fields],
    ]


def _format_guest_worksheet(sheet, headers, last_row):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).column_letter}{max(1, last_row)}"
    sheet.row_dimensions[1].height = 28
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="24735A")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = min(
            max(len(str(header)) + 3, 14), 30
        )


def _guest_export_values(guest, custom_fields):
    return [
        guest.group_serial or "",
        _safe_csv_cell(guest.name),
        _safe_csv_cell(guest.last_name),
        guest.country_code,
        _safe_csv_cell(guest.phone),
        guest.invited_date.isoformat() if guest.invited_date else "",
        _safe_csv_cell(guest.category),
        guest.pax,
        guest.get_guest_type_display() if guest.guest_type else "",
        {
            Guest.SideTag.BRIDE: "Bride Side",
            Guest.SideTag.GROOM: "Groom Side",
        }.get(guest.side_tag, ""),
        guest.get_sub_category_display() if guest.sub_category else "",
        _safe_csv_cell(guest.invite_name),
        _safe_csv_cell(guest.salutation),
        *[
            _safe_csv_cell(guest.custom_data.get(str(field.pk), ""))
            for field in custom_fields
        ],
    ]


@organizer_required
def export_guests(request):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="event-guest-list.csv"'
    response.write("\ufeff")
    writer = csv.writer(response)
    custom_fields = list(GuestField.objects.all())
    writer.writerow(_guest_export_headers(custom_fields))
    for guest in Guest.objects.filter(event=get_active_event(request)).order_by(
        "side_tag", "group_serial", "name", "last_name", "id"
    ):
        writer.writerow(
            [
                _safe_csv_cell(value)
                for value in _guest_export_values(guest, custom_fields)
            ]
        )
    return response


@organizer_required
def export_excel(request):
    workbook = Workbook()
    workbook.remove(workbook.active)
    custom_fields = list(GuestField.objects.all())
    event = get_active_event(request)
    for side, sheet_name in (
        (Guest.SideTag.GROOM, "Groom Side"),
        (Guest.SideTag.BRIDE, "Bride Side"),
    ):
        sheet = workbook.create_sheet(sheet_name)
        headers = _guest_export_headers(custom_fields)
        sheet.append(headers)
        if event:
            for guest in Guest.objects.filter(event=event, side_tag=side).order_by(
                "group_serial", "name", "last_name", "id"
            ):
                sheet.append(_guest_export_values(guest, custom_fields))
        _format_guest_worksheet(sheet, headers, sheet.max_row)
    if event and event.guests.filter(side_tag="").exists():
        sheet = workbook.create_sheet("Unassigned")
        headers = _guest_export_headers(custom_fields)
        sheet.append(headers)
        for guest in event.guests.filter(side_tag="").order_by(
            "group_serial", "name", "last_name", "id"
        ):
            sheet.append(_guest_export_values(guest, custom_fields))
        _format_guest_worksheet(sheet, headers, sheet.max_row)
    content = io.BytesIO()
    workbook.save(content)
    response = HttpResponse(
        content.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="event-guest-list.xlsx"'
    return response


@organizer_required
def guest_template(request):
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["Nimoli Bhandari Events guest workbook"])
    instructions.append(["Use the single Guest List sheet; you do not need to create separate Bride Side and Groom Side sheets."])
    instructions.append(["Use these columns in order: Group ID, First Name, Last Name, Country Code, Phone Number, Invited Dates, Category, Pax, Guest Type, Tag, Sub Category, Invite Name, Salutation."])
    instructions.append(["Keep the header names unchanged; the columns can be moved if needed."])
    instructions.append(["All guest details are optional. Rows with any data are imported; missing Group IDs are assigned automatically."])
    instructions.append(["Choose Bride Side or Groom Side from the Tag dropdown when known; otherwise leave it blank or assign a default side during import."])
    instructions.append(["Reuse one Group ID for everyone in the same family/travel group when known."])
    instructions.append(["Country Code and Phone Number are separate, e.g. +91 and 9876543210."])
    instructions.append(["Guest Type choices: Outstation, Local Stay, Local Non-Stay, International."])
    instructions.append(["Guest details: Invited Dates, Category, Pax, Sub Category, Invite Name, and Salutation."])
    instructions.append(["Extra event-specific columns can be added after the listed fields."])
    instructions.append(["Use the header filter arrows to sort or filter. Import and export are grouped by side, serial number, and guest name."])
    instructions.column_dimensions["A"].width = 115
    for cell in instructions["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    instructions["A1"].font = Font(bold=True, size=16, color="24735A")
    headers = [
        "Group ID",
        "First Name",
        "Last Name",
        "Country Code",
        "Phone Number",
        "Invited Dates",
        "Category",
        "Pax",
        "Guest Type",
        "Tag",
        "Sub Category",
        "Invite Name",
        "Salutation",
    ]
    for field in GuestField.objects.all():
        headers.append(field.label)
    sheet = workbook.create_sheet("Guest List")
    sheet.append(headers)
    _format_guest_worksheet(sheet, headers, 2001)
    sheet.column_dimensions["D"].number_format = "@"
    for row in range(2, 2002):
        sheet.cell(row=row, column=4).number_format = "@"
        sheet.cell(row=row, column=5).number_format = "@"
    guest_type_validation = DataValidation(
        type="list",
        formula1='"Outstation,Local Stay,Local Non-Stay,International"',
        allow_blank=True,
    )
    guest_type_validation.error = "Choose one of the listed guest types."
    guest_type_validation.errorTitle = "Invalid guest type"
    guest_type_validation.showErrorMessage = True
    sheet.add_data_validation(guest_type_validation)
    guest_type_validation.add("I2:I2001")
    side_validation = DataValidation(
        type="list",
        formula1='"Bride Side,Groom Side"',
        allow_blank=True,
    )
    side_validation.error = "Choose Bride Side or Groom Side."
    side_validation.errorTitle = "Invalid guest side"
    side_validation.showErrorMessage = True
    sheet.add_data_validation(side_validation)
    side_validation.add("J2:J2001")

    lists_sheet = workbook.create_sheet("Lists")
    lists_sheet.sheet_state = "hidden"
    lists_sheet.append(["Country Code"])
    for country_code, _ in country_code_choices():
        lists_sheet.append([country_code])
    workbook.defined_names.add(
        DefinedName(
            "CountryCodes",
            attr_text=f"'Lists'!$A$2:$A${lists_sheet.max_row}",
        )
    )
    country_code_validation = DataValidation(
        type="list",
        formula1="=CountryCodes",
        allow_blank=True,
    )
    country_code_validation.error = "Choose a country calling code."
    country_code_validation.errorTitle = "Invalid country code"
    country_code_validation.showErrorMessage = True
    sheet.add_data_validation(country_code_validation)
    country_code_validation.add("D2:D2001")
    content = io.BytesIO()
    workbook.save(content)
    response = HttpResponse(
        content.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="Nimoli-Bhandari-Events-Guest-Template.xlsx"'
    return response


def guest_rsvp(request, token):
    guest = get_object_or_404(Guest, rsvp_token=token)
    custom_fields = list(GuestField.objects.all())
    if request.method == "POST":
        attendance = request.POST.get("attendance")
        if attendance not in {"yes", "no"}:
            messages.error(request, "Choose whether you can attend.")
        elif attendance == "no":
            guest.rsvp_status = Guest.RSVPStatus.DECLINED
            guest.hotel_status = Guest.HotelStatus.NOT_SELECTED
            guest.hotel_check_in = None
            guest.hotel_check_out = None
            guest.room_preference = Guest.RoomPreference.NONE
            guest.save()
            messages.success(request, "Your RSVP has been recorded. Thank you.")
        else:
            room = request.POST.get("room_preference")
            try:
                response_pax = int(request.POST.get("pax", str(guest.pax or 1)))
            except (TypeError, ValueError):
                response_pax = 0
            if room not in Guest.RoomPreference.values:
                messages.error(request, "Choose a valid hotel room preference.")
            elif not 1 <= response_pax <= 100:
                messages.error(request, "Enter a guest count from 1 to 100.")
            else:
                check_in = request.POST.get("hotel_check_in") or None
                check_out = request.POST.get("hotel_check_out") or None
                try:
                    check_in = date.fromisoformat(check_in) if check_in else None
                    check_out = date.fromisoformat(check_out) if check_out else None
                except ValueError:
                    messages.error(request, "Enter valid hotel check-in and check-out dates.")
                else:
                    if (check_in is None) != (check_out is None):
                        messages.error(
                            request,
                            "Enter both hotel dates, or leave both dates blank.",
                        )
                    elif check_in and check_out and check_out <= check_in:
                        messages.error(
                            request, "Check-out must be after the check-in date."
                        )
                    else:
                        custom_data = {}
                        invalid_custom_field = False
                        for field in custom_fields:
                            value = request.POST.get(f"custom_{field.pk}", "").strip()
                            if len(value) > 500:
                                messages.error(
                                    request,
                                    f"“{field.label}” must be 500 characters or fewer.",
                                )
                                invalid_custom_field = True
                                break
                            if value:
                                custom_data[str(field.pk)] = value
                        if not invalid_custom_field:
                            guest.rsvp_status = Guest.RSVPStatus.CONFIRMED
                            guest.room_preference = room
                            guest.pax = response_pax
                            guest.hotel_check_in = check_in
                            guest.hotel_check_out = check_out
                            guest.hotel_status = Guest.HotelStatus.NOT_SELECTED
                            guest.custom_data = custom_data
                            guest.save()
                            messages.success(
                                request,
                                "Your RSVP and stay preferences have been recorded.",
                            )
        return redirect("guest_rsvp", token=guest.rsvp_token)

    response = render(
        request,
        "guests/rsvp.html",
        {
            "guest": guest,
            "room_choices": Guest.RoomPreference.choices,
            "custom_fields": custom_fields,
        },
    )
    response["Cache-Control"] = "no-store"
    response["Referrer-Policy"] = "no-referrer"
    response["X-Robots-Tag"] = "noindex, nofollow, noarchive"
    return response


@require_http_methods(["GET", "POST"])
def organizer_signup(request):
    if (
        not settings.DEBUG
        or request.META.get("REMOTE_ADDR") not in {"127.0.0.1", "::1"}
    ):
        raise Http404

    User = get_user_model()
    if User.objects.filter(is_staff=True).exists():
        return render(
            request,
            "registration/signup.html",
            {"registration_closed": True},
        )

    form = UserCreationForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            if User.objects.filter(is_staff=True).exists():
                return render(
                    request,
                    "registration/signup.html",
                    {"registration_closed": True},
                )
            user = form.save(commit=False)
            user.is_staff = True
            user.save()
        login(request, user)
        messages.success(request, "Your organizer account is ready.")
        return redirect("dashboard")

    return render(
        request,
        "registration/signup.html",
        {"form": form, "registration_closed": False},
    )


def web_manifest(request):
    return JsonResponse(
        {
            "id": "/",
            "name": "Nimoli Bhandari Events",
            "short_name": "Nimoli Bhandari",
            "description": "Manage event RSVPs and hotel coordination.",
            "start_url": "/",
            "scope": "/",
            "display": "standalone",
            "background_color": "#f7f8f5",
            "theme_color": "#24735a",
            "icons": [
                {
                    "src": static("guests/app-icon-192.png"),
                    "sizes": "192x192",
                    "type": "image/png",
                    "purpose": "any",
                },
                {
                    "src": static("guests/app-icon-512.png"),
                    "sizes": "512x512",
                    "type": "image/png",
                    "purpose": "any maskable",
                },
            ],
        },
        content_type="application/manifest+json",
    )


def service_worker(request):
    worker_path = (
        Path(settings.BASE_DIR)
        / "guests"
        / "static"
        / "guests"
        / "service-worker.js"
    )
    response = HttpResponse(
        worker_path.read_text(encoding="utf-8"),
        content_type="application/javascript; charset=utf-8",
    )
    response["Service-Worker-Allowed"] = "/"
    response["Cache-Control"] = "no-cache"
    return response


def offline_page(request):
    return render(request, "guests/offline.html")
